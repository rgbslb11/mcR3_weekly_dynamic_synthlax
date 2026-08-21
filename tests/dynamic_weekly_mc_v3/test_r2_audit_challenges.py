"""Adjudication of the five R2 audit challenges, against the mounted artifacts.

Each test here answers a specific challenge to the convergence by reading the
governed sources rather than by trusting either the builder report or prior
prose.
"""

from pathlib import Path

import pytest
from openpyxl import load_workbook

from ncaaf_engine.simulation.dynamic_weekly_mc_v3 import fcs, postseason, sor, sos, srs
from ncaaf_engine.simulation.dynamic_weekly_mc_v3.config import V3Config
from ncaaf_engine.simulation.dynamic_weekly_mc_v3.errors import GovernanceBlock

ROOT = Path(__file__).resolve().parents[2]
CONFIG = ROOT / "config/dynamic_weekly_mc_v3/v3_experimental.json"
INPUTS = ROOT / "reference/dynamic_weekly_mc_v3/inputs"
CALENDAR = INPUTS / "2026 Playoff Calendar OFFICIAL 2.xlsx"
MODEL_PARAMS = INPUTS / "Model_Parameters_v2_5_APPROVED.xlsx"


# --- Challenge 1: SOR reference Elo ------------------------------------------


def test_the_2026_sor_reference_elo_is_governed_at_1893_3():
    """Not missing governance: it is registered, LOCKED and SOURCE-VERIFIED."""
    wb = load_workbook(MODEL_PARAMS, read_only=True, data_only=True)
    ws = wb["02_PARAMETER_REGISTER"]
    row = next(r for r in ws.iter_rows(values_only=True) if r[0] == "CCG-R_REF")
    assert float(row[4]) == 1893.3          # stored as text in the register
    assert row[7] == "LOCKED"
    assert row[14] == "SOURCE-VERIFIED"
    assert "SOR reference Elo" in str(row[15])
    assert sor.GOVERNED_2026_SOR_R_REF == 1893.3
    assert sor.governed_2026_sor_reference_elo().value == 1893.3


def test_1901_is_explicitly_superseded_in_the_repository():
    """The higher authority the challenge asks for is present and named."""
    wb = load_workbook(MODEL_PARAMS, read_only=True, data_only=True)
    rows = [
        " | ".join("" if v is None else str(v) for v in r)
        for r in wb["20_CANON_MANIFEST_INGEST"].iter_rows(values_only=True)
    ]
    superseded = next(r for r in rows if "R_ref 1901 now stale" in r)
    assert "SUPERSEDED" in superseded
    assert "R-MC-V2" in superseded
    assert "do not use for decisions" in superseded
    change_log = next(r for r in rows if "R_ref 1901->1893.3" in r)
    assert "2026-07-14" in change_log

    with pytest.raises(GovernanceBlock, match="superseded"):
        sor.SorReferenceElo(
            value=1901.0, parameter_id="X", authority="Y", source_artifact="Z"
        )


def test_the_stale_library_default_is_still_refused_separately():
    with pytest.raises(GovernanceBlock, match="stale"):
        sor.SorReferenceElo(
            value=1684.9, parameter_id="X", authority="Y", source_artifact="Z"
        )


def test_remaining_sor_b_items_are_narrower_than_the_reference_elo():
    assert sor.UNRATIFIED_SOR_B_ITEMS == ("P_TO_STRENGTH_TRANSFORM", "REFERENCE_HFA")
    row = sor.weekly_sor_row(
        week=10, as_of="t", team="ARK",
        opponents=[sor.SorOpponent("X", 1600.0, True), sor.SorOpponent("Y", 1500.0, False)],
        reference=sor.governed_2026_sor_reference_elo(), source_hashes={},
        model="m", model_version="v", configuration_version="c",
        observed_at="t", recorded_at="t",
    ).as_dict()
    assert row["r_ref"] == 1893.3
    assert row["provenance"]["unratified_sor_b_items"] == [
        "P_TO_STRENGTH_TRANSFORM", "REFERENCE_HFA"
    ]
    assumptions = " ".join(row["provenance"]["unresolved_assumptions"])
    assert "P_TO_STRENGTH_TRANSFORM" in assumptions
    assert "REFERENCE_HFA" in assumptions
    # R_ref is governed, so it is not among the stamped assumptions.
    assert "R_ref" not in assumptions


def test_the_two_reference_elo_namespaces_stay_separate():
    assert sor.governed_2026_sor_reference_elo().namespace == sor.SOR_REPORT_NAMESPACE
    assert sor.mc_domain_reference_elo().namespace == sor.MC_NAMESPACE
    with pytest.raises(GovernanceBlock, match="keep domains separate"):
        sor.reject_mc_r_ref_overwrite(
            sor.MC_NAMESPACE, sor.governed_2026_sor_reference_elo()
        )


# --- Challenge 2: SRS solver equivalence -------------------------------------


def _league(n_teams: int, n_weeks: int, seed: int) -> list[srs.SrsGame]:
    """Deterministic synthetic league; no RNG state escapes this fixture."""
    games: list[srs.SrsGame] = []
    gid = 0
    teams = [f"T{i:02d}" for i in range(n_teams)]
    for week in range(n_weeks):
        rotated = teams[week:] + teams[:week]
        for i in range(0, len(rotated) - 1, 2):
            gid += 1
            a, b = rotated[i], rotated[i + 1]
            margin = ((gid * seed * 37) % 61) - 30
            games.append(srs.SrsGame(f"G{gid:04d}", a, b, float(margin)))
            games.append(srs.SrsGame(f"G{gid:04d}", b, a, float(-margin)))
    return games


@pytest.mark.parametrize("n_teams,n_weeks", [(8, 3), (12, 5), (24, 7), (40, 9)])
def test_exact_solver_solves_the_governed_system(n_teams, n_weeks):
    games = _league(n_teams, n_weeks, seed=3)
    report = srs.assert_solver_equivalence(games)
    assert report["worst_residual"] < 1e-9
    assert report["worst_component_centring"] < 1e-9


@pytest.mark.parametrize("n_teams,n_weeks", [(8, 3), (12, 5), (24, 7), (40, 9)])
def test_exact_and_converged_iterative_agree_on_values_and_ordering(n_teams, n_weeks):
    games = _league(n_teams, n_weeks, seed=5)
    report = srs.assert_solver_equivalence(games)
    assert report["iterative_converged"] is True
    assert report["max_abs_difference"] < 1e-9
    # Ordering is only a meaningful comparison where no two ratings sit closer
    # together than the agreement tolerance; at a near-exact tie either order is
    # equally correct and the name tiebreak decides.
    if report["min_rating_gap"] > report["max_abs_difference"]:
        assert report["orderings_match"] is True


def test_centering_is_preserved():
    games = _league(24, 7, seed=11)
    ratings = srs.compute_srs(games)
    assert sum(ratings.values()) == pytest.approx(0.0, abs=1e-9)


def test_the_plus_minus_24_cap_is_preserved():
    assert srs.SRS_MARGIN_CAP == 24.0
    capped = [
        srs.SrsGame("g1", "A", "B", 24.0), srs.SrsGame("g1", "B", "A", -24.0),
        srs.SrsGame("g2", "A", "C", 7.0), srs.SrsGame("g2", "C", "A", -7.0),
    ]
    blown = [
        srs.SrsGame("g1", "A", "B", 300.0), srs.SrsGame("g1", "B", "A", -300.0),
        srs.SrsGame("g2", "A", "C", 7.0), srs.SrsGame("g2", "C", "A", -7.0),
    ]
    assert srs.compute_srs(blown) == srs.compute_srs(capped)


def test_no_early_week_disconnected_graph_instability():
    """Week 1 is three isolated pairs. It must solve, not raise or oscillate."""
    games: list[srs.SrsGame] = []
    for i, (a, b, m) in enumerate([("A", "B", 21.0), ("C", "D", -3.0), ("E", "F", 10.0)], 1):
        games += [srs.SrsGame(f"W{i}", a, b, m), srs.SrsGame(f"W{i}", b, a, -m)]
    components = srs.srs_components(games)
    assert len(components) == 3
    ratings = srs.compute_srs(games)
    for component in components:
        assert sum(ratings[t] for t in component) == pytest.approx(0.0, abs=1e-12)
    assert srs.assert_solver_equivalence(games)["worst_residual"] < 1e-9


def test_the_replaced_jacobi_scheme_returned_a_wrong_ordering():
    """Why the solver was changed: the old scheme did not solve the system."""
    games = [
        srs.SrsGame("g1", "A", "B", 30.0), srs.SrsGame("g1", "B", "A", -30.0),
        srs.SrsGame("g2", "A", "C", 7.0), srs.SrsGame("g2", "C", "A", -7.0),
    ]
    exact = srs.compute_srs(games)
    assert srs.srs_ordering(exact) == ["A", "C", "B"]
    assert max(abs(v) for v in srs.srs_residuals(games, exact).values()) < 1e-9

    # The oscillating fixed point the Jacobi sweep settles on.
    jacobi = {"A": 0.0, "B": -8.5, "C": 8.5}
    assert srs.srs_ordering(jacobi) == ["C", "A", "B"]
    assert max(abs(v) for v in srs.srs_residuals(games, jacobi).values()) > 1.0


def test_srs_is_not_claimed_canonically_validated():
    """No canonical spec, implementation or anchor is mounted, so nothing claims it."""
    assert srs.CANONICAL_SRS_SPEC_MOUNTED is False
    assert srs.CANONICAL_VALIDATION_ANCHORS == ()
    assert srs.CANONICAL_SRS_REFERENCE_IMPLEMENTATION is None
    with pytest.raises(GovernanceBlock, match=srs.CANONICAL_VALIDATION_BLOCKER):
        srs.require_canonical_validated_srs()
    witness = srs.witness_as_dict(srs.compute_srs(_league(8, 3, seed=2)))
    assert witness["canonical_validation_status"] == "NOT_VALIDATED_AGAINST_CANONICAL_ANCHORS"


def test_srs_over_40_semantics_are_undefined_and_fail_closed():
    assert srs.SRS_OVER_40_SEMANTICS is None
    with pytest.raises(GovernanceBlock, match="not defined in any mounted artifact"):
        srs.require_srs_over_40(41.0)


# --- Challenge 4: quarterfinal source mapping --------------------------------


def test_the_calendar_workbook_has_no_formulas_named_ranges_or_hidden_content():
    wb = load_workbook(CALENDAR, data_only=False)
    assert wb.sheetnames == [
        "Playoff_Calendar", "Bracket_Flow", "Rest_Integrity", "Rulings_Register", "Open_Items"
    ]
    assert list(wb.defined_names) == []
    for name in wb.sheetnames:
        ws = wb[name]
        assert ws.sheet_state == "visible"
        assert not [r for r, d in ws.row_dimensions.items() if d.hidden]
        assert not [c for c, d in ws.column_dimensions.items() if d.hidden]
        assert not list(ws.merged_cells.ranges)
        formulas = [
            c.coordinate for row in ws.iter_rows() for c in row
            if isinstance(c.value, str) and c.value.startswith("=")
        ]
        # The only formulas in the workbook are date arithmetic on Rest_Integrity.
        assert formulas == (["D3", "D4", "D5"] if name == "Rest_Integrity" else [])


def test_the_quarterfinal_rows_never_name_a_specific_first_round_game():
    """The exact source text, at the exact coordinates."""
    wb = load_workbook(CALENDAR, data_only=True)
    calendar = wb["Playoff_Calendar"]
    assert calendar["C12"].value == "Seed 1 v W(R1)"
    assert calendar["C14"].value == "Seed 2 v W(R1)"
    assert calendar["C15"].value == "Seed 3 v W(R1)"
    assert calendar["C13"].value == "Seed 4 v W(R1)"
    # The four first-round rows share one composition string, so they are not even
    # individually identified by seed pairing in this workbook, and none of their
    # notes names a destination quarterfinal slot.
    for coord in ("C8", "C9", "C10", "C11"):
        assert calendar[coord].value == "seeds 5–10 / PI winners"
    assert calendar["H8"].value == "→ QF"
    assert calendar["H9"].value == "→ QF"
    for coord in ("H8", "H9", "H10", "H11"):
        note = str(calendar[coord].value)
        for slot in ("E", "F", "G", "H"):
            assert f"GAME {slot}" not in note.upper()
            assert f"QF {slot}" not in note.upper()
        assert "→ QF" in note or "P-7" in note

    flow = wb["Bracket_Flow"]
    assert flow["B6"].value == "E = 1 v W(R1); F = 2 v W(R1); G = 3 v W(R1); H = 4 v W(R1)"
    assert "W(R1)" in flow["B6"].value
    for game in ("A", "B", "C", "D"):
        assert f"W({game})" not in flow["B6"].value


def test_the_source_still_fails_to_bind_the_edges_and_is_not_reinterpreted():
    """R2-A's finding stands: the workbook never bound the four R1 winners.

    The R3 ruling fills that gap by successor Chairman authority. It does not
    make the workbook say something it never said, so this flag stays False.
    """
    assert postseason.QUARTERFINAL_SLOT_EDGES_STATED_IN_ARTIFACT is False
    assert postseason.QUARTERFINAL_MAPPING_RESOLUTION_REASON == (
        "SUCCESSOR_DIRECT_CHAIRMAN_AUTHORITY"
    )
    assert postseason.QUARTERFINAL_MAPPING_RESOLUTION_REASON != (
        "SOURCE_WORKBOOK_CONTAINED_MAPPING"
    )


# --- Challenge 5: FCS scale, not FCS policy ----------------------------------


def test_the_fcs_rating_policy_is_not_reopened():
    assert fcs.FCS_RATING_POLICY_RESOLVED is True
    assert fcs.require_governed_fcs_policy("FIXED_ELO_1250").fixed_elo == 1250.0


def test_the_missing_bridge_is_classified_as_a_model_scale_adapter():
    assert fcs.FCS_UNIFIED_SCALE_BLOCKER.startswith("model_scale.")
    assert "GOVERNED" not in fcs.FCS_UNIFIED_SCALE_BLOCKER
    with pytest.raises(GovernanceBlock, match="rating policy is settled"):
        fcs.require_fcs_unified_points()


def test_the_simulation_layer_genuinely_requires_the_bridge():
    from ncaaf_engine.simulation.dynamic_weekly_mc_v3.engine import DynamicWeeklyMCV3
    from ncaaf_engine.simulation.dynamic_weekly_mc_v3.inputs import load_teams

    cfg = V3Config.from_json(CONFIG)
    teams = load_teams(cfg.inputs.canonical_master_md, cfg.inputs.unified_preseason_ratings_xlsx)
    fcs_teams = {s: t for s, t in teams.items() if t.entity_scope == "SCHEDULE_ONLY_FCS"}
    assert len(fcs_teams) == 13
    assert all(t.preseason_strength_points is None for t in fcs_teams.values())
    with pytest.raises(GovernanceBlock, match="no unified preseason strength points"):
        DynamicWeeklyMCV3(cfg)._initialize_states(teams)


def test_the_superseded_v2_1_bridge_is_recorded():
    wb = load_workbook(
        INPUTS / "V2_1_STATIC_CONTROL_2026_CFB_10000_Season_Monte_Carlo.xlsx",
        read_only=True, data_only=True,
    )
    text = " ".join(
        str(c) for row in wb["Methodology"].iter_rows(values_only=True) for c in row if c
    )
    assert "translated from Board I-H equivalent to unified points" in text
    assert "Board I-H equivalent" in fcs.V2_1_FCS_BRIDGE


# --- Challenge 6: OWP / OOWP semantics ---------------------------------------


def test_all_six_semantics_questions_are_stated_in_the_required_ruling():
    required = sos.required_ruling_text()
    assert len(sos.REQUIRED_SEMANTICS_RULINGS) == 6
    for fragment in (
        "removed from that opponent's record",
        "team-averaged or schedule-instance-weighted",
        "repeated opponent",
        "How is OOWP constructed",
        "schedule-only FCS opponent records",
        "zero qualifying games",
    ):
        assert fragment in required


def test_exactly_one_narrow_blocker_covers_all_six():
    with pytest.raises(GovernanceBlock, match=sos.SOS_SEMANTICS_BLOCKER) as excinfo:
        sos.require_governed_sos_semantics(None)
    message = str(excinfo.value)
    assert message.count(sos.SOS_SEMANTICS_BLOCKER) == 1
    for i in range(1, 7):
        assert f"{i}." in message


def test_the_nearest_authority_answers_none_of_the_six():
    """V2.1's chain names an OWP-like quantity and defines nothing."""
    wb = load_workbook(
        INPUTS / "V2_1_STATIC_CONTROL_2026_CFB_10000_Season_Monte_Carlo.xlsx",
        read_only=True, data_only=True,
    )
    rows = [
        " | ".join("" if v is None else str(v) for v in r)
        for r in wb["Methodology"].iter_rows(values_only=True)
    ]
    chain = next(r for r in rows if "average opponents winning percentage" in r)
    assert "Committee method A" in chain
    joined = " ".join(rows).lower()
    assert "oowp" not in joined
    # None of the six semantics questions is addressed. ("excluded" does appear,
    # but only about the W16 Army-Navy game and CFP selection, not about OWP.)
    for undefined in (
        "opponent's games against",
        "denominator",
        "instance-weight",
        "schedule-instance",
        "zero qualifying",
        "repeated opponent",
    ):
        assert undefined not in joined
    exclusion_context = [r for r in rows if "exclud" in r.lower()]
    assert all("CFP selection" in r for r in exclusion_context)


def test_the_ruling_must_answer_the_oowp_construction_question():
    """The two constructions genuinely differ on unequal opponent game counts."""
    rows = []
    n = [0]

    def game(week, winner, loser):
        n[0] += 1
        gid = f"F{n[0]:04d}"
        rows.append(sos.GameResult(gid, week, winner, loser, True))
        rows.append(sos.GameResult(gid, week, loser, winner, False))

    # A's two opponents play very different numbers of games, which is exactly
    # when pooling every opponent-of-opponent diverges from averaging per opponent.
    game(1, "A", "B")
    game(2, "A", "C")
    game(3, "B", "D")
    game(4, "B", "E")
    game(5, "B", "F")
    game(6, "C", "G")
    game(7, "D", "E")
    game(8, "F", "G")
    ledger = sos.ResumeLedger(rows)
    base = dict(
        semantics_id="F", authority="TEST_FIXTURE", exclude_rated_team_from_owp=True,
        instance_weighting="PER_GAME", exclude_rated_team_from_oowp=True,
    )
    per_opponent = sos.SosSemantics(**base, oowp_construction="MEAN_OF_OPPONENT_OWP")
    pooled = sos.SosSemantics(**base, oowp_construction="MEAN_OVER_ALL_OPPONENT_OPPONENTS")
    assert sos.opponent_opponent_win_pct(ledger, "A", per_opponent) != (
        sos.opponent_opponent_win_pct(ledger, "A", pooled)
    )


def test_zero_qualifying_games_policy_is_explicit_not_implicit():
    ledger = sos.ResumeLedger([sos.GameResult("G1", 1, "A", "B", True),
                               sos.GameResult("G1", 1, "B", "A", False)])
    blocking = sos.SosSemantics(
        semantics_id="F", authority="TEST_FIXTURE", exclude_rated_team_from_owp=True,
        instance_weighting="PER_GAME", exclude_rated_team_from_oowp=True,
        zero_qualifying_games="BLOCK",
    )
    # A carries one game; excluding its only opponent leaves zero qualifying games.
    with pytest.raises(GovernanceBlock, match="zero qualifying games"):
        sos.win_pct(ledger, "A", excluding_opponent="B", semantics=blocking)


def test_schedule_only_fcs_treatment_changes_owp():
    rows = []
    n = [0]

    def game(week, winner, loser):
        n[0] += 1
        gid = f"F{n[0]:04d}"
        rows.append(sos.GameResult(gid, week, winner, loser, True))
        rows.append(sos.GameResult(gid, week, loser, winner, False))

    # The FCS opponent has a losing record and the FBS opponent a winning one,
    # so dropping the FCS entity moves A's OWP.
    game(1, "A", "FCS1")
    game(2, "A", "B")
    game(3, "B", "C")
    game(4, "B", "D")
    game(5, "FCS2", "FCS1")
    game(6, "FCS3", "FCS1")
    ledger = sos.ResumeLedger(rows)
    base = dict(
        semantics_id="F", authority="TEST_FIXTURE", exclude_rated_team_from_owp=True,
        instance_weighting="PER_GAME", exclude_rated_team_from_oowp=True,
        schedule_only_fcs_ids=frozenset({"FCS1", "FCS2"}),
    )
    included = sos.SosSemantics(**base, schedule_only_fcs_treatment="INCLUDE")
    excluded = sos.SosSemantics(**base, schedule_only_fcs_treatment="EXCLUDE_FROM_ALL")
    assert sos.opponent_win_pct(ledger, "A", included) != (
        sos.opponent_win_pct(ledger, "A", excluded)
    )
