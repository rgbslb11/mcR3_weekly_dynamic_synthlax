"""Lane F1: the FCS Elo-1250 -> V3 unified-points model-scale adapter.

The FCS *rating policy* is settled by ruling R2-FCS-ELO-1250 and is not under
test here. What is under test is the model-scale adapter the ruling deliberately
leaves open, and the claim this lane makes about it: that no governed Elo ->
unified-neutral-points mapping exists in the mounted artifacts, and therefore
that `model_scale.FCS_ELO_1250_TO_V3_POINT_SCALE_ADAPTER` is correctly retained.

That claim is only worth as much as its reproduction, so the two transforms that
*do* exist are recomputed here from the workbooks rather than quoted.
"""

from pathlib import Path

import pytest
from openpyxl import load_workbook

from ncaaf_engine.simulation.dynamic_weekly_mc_v3 import blocker_report as br
from ncaaf_engine.simulation.dynamic_weekly_mc_v3 import fcs, fcs_scale
from ncaaf_engine.simulation.dynamic_weekly_mc_v3.config import V3Config
from ncaaf_engine.simulation.dynamic_weekly_mc_v3.engine import DynamicWeeklyMCV3
from ncaaf_engine.simulation.dynamic_weekly_mc_v3.errors import (
    GovernanceBlock,
    InputValidationError,
)
from ncaaf_engine.simulation.dynamic_weekly_mc_v3.inputs import load_schedule, load_teams

ROOT = Path(__file__).resolve().parents[2]
CONFIG = ROOT / "config/dynamic_weekly_mc_v3/v3_experimental.json"
INPUTS = ROOT / "reference/dynamic_weekly_mc_v3/inputs"
EXPERIMENTAL = (
    ROOT / "config/dynamic_weekly_mc_v3/experimental/fcs_point_scale_regimes.json"
)
CALIBRATION_REGIMES = (
    ROOT / "config/dynamic_weekly_mc_v3/experimental/calibration_regimes.json"
)


@pytest.fixture(scope="module")
def cfg():
    return V3Config.from_json(CONFIG)


@pytest.fixture(scope="module")
def teams(cfg):
    return load_teams(cfg.inputs.canonical_master_md, cfg.inputs.unified_preseason_ratings_xlsx)


@pytest.fixture(scope="module")
def schedule(cfg):
    return load_schedule(cfg.inputs.schedule_xlsx)


@pytest.fixture(scope="module")
def coverage(teams, schedule):
    return fcs_scale.fcs_scale_coverage(teams, schedule)


# --- step 1: reproduce the two transforms that do exist ----------------------


def test_elo_board_transform_reproduces_on_every_board_rated_row():
    """The Build Manifest's own residual claim, recomputed rather than quoted."""
    wb = load_workbook(
        INPUTS / "POWER_CRUNCH_2026_Team_Reconciled_Master_134_v2_2.xlsx",
        read_only=True,
        data_only=True,
    )
    rows = list(wb["Reconciled Master"].iter_rows(values_only=True))
    idx = {name: i for i, name in enumerate(rows[0])}

    residuals = []
    for row in rows[1:]:
        if row[idx["schedule_id"]] is None:
            continue
        try:
            board = float(row[idx["board_power_H"]])
        except (TypeError, ValueError):
            continue  # UNRATED — the FCS rows, handled below
        elo = float(row[idx["scheme_master_primary_elo"]])
        residuals.append(abs(fcs_scale.elo_from_board_power_h(board) - elo))

    assert len(residuals) == fcs_scale.ELO_BOARD_RECORDED_ROW_COUNT == 121
    # The manifest states the residual to four places; the reproduction is kept
    # at full precision so the rounding is visible rather than hidden in a bound.
    assert max(residuals) == pytest.approx(
        fcs_scale.ELO_BOARD_REPRODUCED_MAX_RESIDUAL, abs=1e-15
    )
    assert round(max(residuals), 4) == fcs_scale.ELO_BOARD_RECORDED_MAX_RESIDUAL


def test_unified_neutral_field_points_reproduces_exactly_on_all_121_fbs_rows():
    """`points = 14 * UnifiedMasterZ` over the four-family ensemble, to 0.0 error."""
    wb = load_workbook(
        INPUTS / "2026_CFB_2026_Preseason_Unified_Power_Ratings.xlsx",
        read_only=True,
        data_only=True,
    )
    rows = list(wb["Master Ratings"].iter_rows(min_row=4, values_only=True))
    idx = {name: i for i, name in enumerate(rows[0]) if name}
    data = [r for r in rows[1:] if r[idx["Team"]]]
    assert len(data) == fcs_scale.UNIFIED_RATINGS_FBS_ROWS == 121

    for row in data:
        master_z = fcs_scale.unified_master_z(
            trueskill_z=fcs_scale.standardize(
                row[idx["TrueSkill 2026 Preseason μ"]], "trueskill"
            ),
            litkenhous_z=fcs_scale.standardize(
                row[idx["Litkenhous Adjusted Power"]], "litkenhous"
            ),
            pure_baxter_z=fcs_scale.standardize(
                row[idx["2026 Pure Baxter Rating"]], "pure_baxter"
            ),
            board_i_h_z=fcs_scale.standardize(
                row[idx["Board I-H Power Rating (power_H)"]], "board_i_h"
            ),
            board_j_b_z=fcs_scale.standardize(
                row[idx["Board J-B Run2 Power Rating (power_Run2)"]], "board_j_b"
            ),
        )
        assert master_z == pytest.approx(float(row[idx["Unified Master Z"]]), abs=1e-12)
        assert fcs_scale.unified_neutral_field_points(master_z) == pytest.approx(
            float(row[idx["Unified Neutral-Field Points"]]), abs=1e-12
        )
        assert fcs_scale.unified_master_power_index(master_z) == pytest.approx(
            float(row[idx["Unified Master Power Index"]]), abs=1e-12
        )


def test_the_observed_fbs_points_span_is_the_reproduced_one():
    wb = load_workbook(
        INPUTS / "2026_CFB_2026_Preseason_Unified_Power_Ratings.xlsx",
        read_only=True,
        data_only=True,
    )
    rows = list(wb["Master Ratings"].iter_rows(min_row=4, values_only=True))
    idx = {name: i for i, name in enumerate(rows[0]) if name}
    points = [
        float(r[idx["Unified Neutral-Field Points"]]) for r in rows[1:] if r[idx["Team"]]
    ]
    assert min(points) == pytest.approx(fcs_scale.OBSERVED_FBS_UNIFIED_POINTS_MIN, abs=5e-5)
    assert max(points) == pytest.approx(fcs_scale.OBSERVED_FBS_UNIFIED_POINTS_MAX, abs=5e-5)


def test_the_ratings_workbook_holds_no_fcs_row_at_all():
    """The reason the ensemble cannot be run for an FCS entity, proven not asserted."""
    wb = load_workbook(
        INPUTS / "2026_CFB_2026_Preseason_Unified_Power_Ratings.xlsx",
        read_only=True,
        data_only=True,
    )
    rows = list(wb["Master Ratings"].iter_rows(min_row=4, values_only=True))
    idx = {name: i for i, name in enumerate(rows[0]) if name}
    sids = {str(r[idx["Schedule ID"]]) for r in rows[1:] if r[idx["Team"]]}
    assert len(sids) == 121
    assert not (sids & fcs_scale.SCHEDULE_ONLY_FCS_IDS)


def test_every_fcs_row_is_board_unrated_with_the_equivalent_recorded_not_issued():
    wb = load_workbook(
        INPUTS / "POWER_CRUNCH_2026_Team_Reconciled_Master_134_v2_2.xlsx",
        read_only=True,
        data_only=True,
    )
    rows = list(wb["Reconciled Master"].iter_rows(values_only=True))
    idx = {name: i for i, name in enumerate(rows[0])}
    fcs_rows = [r for r in rows[1:] if r[idx["entity_scope"]] == "SCHEDULE_ONLY_FCS"]

    assert len(fcs_rows) == 13
    assert {str(r[idx["schedule_id"]]) for r in fcs_rows} == fcs_scale.SCHEDULE_ONLY_FCS_IDS
    assert {str(r[idx["board_power_H"]]) for r in fcs_rows} == {"UNRATED"}
    assert {float(r[idx["sim_rating"]]) for r in fcs_rows} == {
        fcs_scale.SUPERSEDED_OPERATOR_COMPOSITE_ELO
    }
    assert {float(r[idx["board_power_H_equivalent"]]) for r in fcs_rows} == {
        fcs_scale.RECORDED_NOT_ISSUED_BOARD_EQUIVALENT
    }


def test_the_recorded_board_equivalent_is_the_inverted_transform():
    """Why the ruling closes the route rather than merely disliking the number."""
    assert round(fcs_scale.DERIVED_INVERSE_OF_SUPERSEDED_COMPOSITE, 6) == (
        fcs_scale.RECORDED_NOT_ISSUED_BOARD_EQUIVALENT
    )


# --- step 1 conclusion: no governed mapping exists ---------------------------


def test_no_governed_elo_to_unified_points_mapping_exists():
    result = fcs_scale.search_governed_elo_to_points_mapping()
    assert result.outcome == fcs_scale.NO_GOVERNED_MAPPING
    assert result.governed_mapping_exists is False
    assert len(result.reproducible_mappings) == 2
    assert len(result.reasons_not_composable) >= 3
    assert len(result.artifacts_searched) == 8


def test_board_i_h_is_one_eighth_of_the_composite_and_three_families_are_missing():
    """The arithmetic reason the two transforms do not compose."""
    board_i_h_share = (
        fcs_scale.FAMILY_WEIGHTS["board_family"]
        * fcs_scale.BOARD_FAMILY_MEMBER_WEIGHTS["board_i_h"]
    )
    assert board_i_h_share == pytest.approx(0.125)
    assert sum(fcs_scale.FAMILY_WEIGHTS.values()) == pytest.approx(1.0)
    assert set(fcs_scale.REQUIRED_RATING_FAMILIES) == set(fcs_scale.FAMILY_WEIGHTS)


# --- the 13 schedule-only FCS identities -------------------------------------


def test_the_thirteen_identities_are_exactly_the_canonical_set(teams, coverage):
    live = {sid for sid, t in teams.items() if t.entity_scope == "SCHEDULE_ONLY_FCS"}
    assert live == fcs_scale.SCHEDULE_ONLY_FCS_IDS
    assert len(live) == 13
    assert coverage["fcs_identity_count"] == 13
    assert coverage["identities_resolvable"] == 0


@pytest.mark.parametrize("schedule_id", sorted(fcs_scale.SCHEDULE_ONLY_FCS_IDS))
def test_each_fcs_identity_is_unresolvable_on_the_points_axis(schedule_id, teams, coverage):
    """All 13, individually: no points, no HFA modifier, adapter refuses."""
    team = teams[schedule_id]
    assert team.entity_scope == "SCHEDULE_ONLY_FCS"
    assert team.preseason_strength_points is None
    assert team.preseason_power_index is None
    assert team.hfa_modifier is None

    row = next(r for r in coverage["fcs_identities"] if r["schedule_id"] == schedule_id)
    assert row["governed_elo"] == fcs.FCS_FIXED_ELO == 1250.0
    assert row["unified_points_resolvable"] is False
    assert row["blocked_by"] == fcs.FCS_UNIFIED_SCALE_BLOCKER

    with pytest.raises(GovernanceBlock, match="rating policy is settled"):
        fcs_scale.GOVERNED_FCS_POINT_SCALE_ADAPTER.unified_points()


# --- every FBS-versus-FCS schedule game --------------------------------------


def test_the_schedule_contains_exactly_the_expected_fbs_v_fcs_games(coverage):
    assert coverage["fbs_v_fcs_game_count"] == 15
    assert tuple(g["game_id"] for g in coverage["fbs_v_fcs_games"]) == (
        fcs_scale.FBS_V_FCS_GAME_IDS
    )
    assert coverage["weeks_affected"] == list(fcs_scale.FBS_V_FCS_WEEKS) == [2, 3, 4, 5, 12]
    assert coverage["games_simulatable"] == 0


@pytest.mark.parametrize("game_id", fcs_scale.FBS_V_FCS_GAME_IDS)
def test_each_fbs_v_fcs_game_is_blocked_by_the_missing_adapter(game_id, coverage):
    row = next(g for g in coverage["fbs_v_fcs_games"] if g["game_id"] == game_id)
    assert row["game_type"] == "REG"
    assert row["requires_points_adapter"] is True
    assert row["simulatable"] is False
    assert row["blocked_by"] == fcs.FCS_UNIFIED_SCALE_BLOCKER
    assert (game_id in fcs_scale.FCS_HOME_GAME_IDS) is row["fcs_entity_is_home"]
    assert row["requires_home_hfa_modifier"] is row["fcs_entity_is_home"]


def test_the_three_fcs_hosted_games_also_need_a_governed_hfa_modifier(coverage):
    """A second gap: SCHED-HFA-BASE is `modifier * 3.5` and the modifier is None."""
    assert coverage["fcs_home_game_ids"] == list(fcs_scale.FCS_HOME_GAME_IDS)
    assert len(fcs_scale.FCS_HOME_GAME_IDS) == 3
    with pytest.raises(GovernanceBlock, match="do not default it to 1.0"):
        fcs_scale.GOVERNED_FCS_POINT_SCALE_ADAPTER.home_hfa_modifier()


def test_every_fcs_game_is_flagged_as_such_in_the_governed_schedule(schedule, coverage):
    """The engine's own fcs_game flag agrees with the adapter's coverage set."""
    flagged = {g.game_id for g in schedule if g.fcs_game}
    assert flagged == {g["game_id"] for g in coverage["fbs_v_fcs_games"]}


# --- the refusal surface ------------------------------------------------------


@pytest.mark.parametrize("value", [0.294, 0.297, 0.297514])
def test_the_recorded_board_equivalents_stay_refused(value):
    with pytest.raises(GovernanceBlock, match="conversion rule"):
        fcs_scale.reject_manufactured_board_equivalent(value)


def test_the_board_equivalent_newly_computable_from_elo_1250_is_refused_by_name():
    """The ruling changed the Elo; it did not reopen the inversion."""
    manufactured = fcs_scale.MANUFACTURED_BOARD_EQUIVALENT_FOR_ELO_1250
    assert manufactured == pytest.approx(0.15000218648009245, abs=1e-15)
    with pytest.raises(GovernanceBlock, match="inverted POWER_CRUNCH Elo/Board transform"):
        fcs_scale.reject_manufactured_board_equivalent(manufactured)


@pytest.mark.parametrize("value", [1250.0, 1397.51, 1893.3, 65.0, 68.0])
def test_an_elo_layer_constant_is_refused_as_a_football_point_value(value):
    with pytest.raises(GovernanceBlock, match="Elo axis"):
        fcs_scale.reject_elo_read_as_points(value)


def test_inverting_the_transform_remains_refused_and_is_never_reimplemented():
    with pytest.raises(GovernanceBlock, match="forbidden"):
        fcs.invert_board_transform(fcs.FCS_FIXED_ELO)

    # The adapter exposes no Elo -> Board callable at all, and the inverse
    # expression appears exactly twice in its source: the two DERIVED constants
    # that exist solely to be refused by name.
    assert not hasattr(fcs_scale, "invert_board_transform")
    assert not any(
        name.startswith("board_power_h_from") or name.startswith("invert")
        for name in dir(fcs_scale)
    )
    source = Path(fcs_scale.__file__).read_text(encoding="utf-8")
    assert source.count(") / ELO_BOARD_SLOPE") == 2
    for constant in (
        fcs_scale.DERIVED_INVERSE_OF_SUPERSEDED_COMPOSITE,
        fcs_scale.MANUFACTURED_BOARD_EQUIVALENT_FOR_ELO_1250,
    ):
        with pytest.raises(GovernanceBlock):
            fcs_scale.reject_manufactured_board_equivalent(constant)


def test_the_forward_transform_direction_is_still_available():
    """Board -> Elo is the direction the manifest issued; it is not forbidden."""
    assert fcs_scale.elo_from_board_power_h(0.757511) == pytest.approx(1857.51, abs=0.015)


def test_an_ungoverned_fcs_elo_is_refused_by_the_adapter():
    for bad in (1300.0, 1397.51, 0.0):
        with pytest.raises(GovernanceBlock, match="not the governed value"):
            fcs_scale.require_governed_fcs_elo(bad)
    assert fcs_scale.require_governed_fcs_elo(1250.0) == 1250.0


def test_a_candidate_outside_the_observed_span_is_reported_not_refused():
    """A reviewer signal, deliberately not an invented validation band."""
    assert fcs_scale.outside_observed_fbs_points_span(-40.0) is True
    assert fcs_scale.outside_observed_fbs_points_span(-10.0) is False


# --- the adapter interface ----------------------------------------------------


def test_the_governed_adapter_holds_no_value_and_fails_closed():
    adapter = fcs_scale.GOVERNED_FCS_POINT_SCALE_ADAPTER
    assert fcs_scale.CANONICAL_FCS_UNIFIED_POINTS is None
    assert fcs_scale.CANONICAL_FCS_HOME_HFA_MODIFIER is None
    assert adapter.authority == fcs_scale.ADAPTER_AUTHORITY_GOVERNED
    assert adapter.governed_for_execution is False
    with pytest.raises(GovernanceBlock, match=fcs.FCS_UNIFIED_SCALE_BLOCKER):
        adapter.unified_points()


def test_the_governed_adapter_is_refused_at_the_engine_seam_while_it_holds_no_rule():
    with pytest.raises(GovernanceBlock, match="holds no governed"):
        fcs_scale.require_governed_adapter(fcs_scale.GOVERNED_FCS_POINT_SCALE_ADAPTER)


def test_an_experimental_adapter_can_never_reach_a_governed_run():
    regime = fcs_scale.CandidateFcsScaleRegime(
        regime_id="X1",
        values={"fcs_unified_points_equivalent": -25.0, "fcs_home_hfa_modifier": 1.0},
        rationale="test only",
    )
    adapter = fcs_scale.ExperimentalFcsPointScaleAdapter(regime)
    assert adapter.governed_for_execution is False
    assert adapter.unified_points() == -25.0
    assert adapter.home_hfa_modifier() == 1.0
    assert adapter.as_dict()["status"] == "EXPERIMENTAL_CANDIDATE_NOT_PROMOTED"
    with pytest.raises(GovernanceBlock, match="may not be consumed"):
        fcs_scale.require_governed_adapter(adapter)


def test_the_adapter_interface_has_no_default_implementation_returning_a_number():
    with pytest.raises(TypeError):
        fcs_scale.FcsPointScaleAdapter()  # type: ignore[abstract]


# --- the experimental harness -------------------------------------------------


def test_the_harness_ships_with_zero_candidate_regimes():
    """Authoring one would invent the calibration the blocker records as missing."""
    assert fcs_scale.load_candidate_fcs_scale_regimes(EXPERIMENTAL) == []
    assert fcs_scale.fcs_point_scale_status()["experimental_regimes_authored"] == 0


def test_regimes_cannot_be_loaded_from_the_canonical_config():
    with pytest.raises(GovernanceBlock, match="canonical"):
        fcs_scale.load_candidate_fcs_scale_regimes(CONFIG)


def test_regimes_cannot_be_loaded_from_a_non_experimental_path(tmp_path):
    stray = tmp_path / "fcs_point_scale_regimes.json"
    stray.write_text('{"regimes": []}', encoding="utf-8")
    with pytest.raises(GovernanceBlock, match="experimental config path"):
        fcs_scale.load_candidate_fcs_scale_regimes(stray)


def test_a_regime_may_not_set_fields_outside_its_own_namespace():
    with pytest.raises(InputValidationError, match="unknown fields"):
        fcs_scale.CandidateFcsScaleRegime(
            regime_id="X", values={"game_sd_points": 17.0}, rationale=""
        )


def test_a_regime_may_not_be_authored_as_anything_other_than_experimental():
    with pytest.raises(GovernanceBlock, match="must be EXPERIMENTAL"):
        fcs_scale.CandidateFcsScaleRegime(
            regime_id="X", values={}, rationale="", status="CANONICAL"
        )


@pytest.mark.parametrize(
    "value", [1250.0, 1397.51, 0.297514, fcs_scale.MANUFACTURED_BOARD_EQUIVALENT_FOR_ELO_1250]
)
def test_a_closed_route_is_refused_at_the_moment_a_candidate_is_authored(value):
    with pytest.raises(GovernanceBlock):
        fcs_scale.CandidateFcsScaleRegime(
            regime_id="X", values={"fcs_unified_points_equivalent": value}, rationale=""
        )


def test_a_non_finite_candidate_is_refused():
    with pytest.raises(InputValidationError, match="non-finite"):
        fcs_scale.CandidateFcsScaleRegime(
            regime_id="X", values={"fcs_unified_points_equivalent": float("inf")}, rationale=""
        )


def test_no_objective_means_no_best_regime():
    with pytest.raises(GovernanceBlock, match="No evaluation objective"):
        fcs_scale.require_fcs_scale_objective(None)


# --- the promotion gate -------------------------------------------------------


@pytest.fixture()
def candidate():
    return fcs_scale.CandidateFcsScaleRegime(
        regime_id="FCS-SCALE-CANDIDATE-TEST",
        values={"fcs_unified_points_equivalent": -25.0},
        rationale="test only",
    )


def test_promotion_without_an_authority_is_refused(candidate):
    with pytest.raises(GovernanceBlock, match="no promotion authority"):
        fcs_scale.promote_fcs_scale_regime(candidate)


def test_winning_an_evaluation_is_not_authority_to_be_promoted(candidate):
    with pytest.raises(GovernanceBlock, match="A top ranking is not an authority"):
        fcs_scale.promote_fcs_scale_regime(candidate, ranked_first=True)


def test_promotion_without_a_human_token_is_refused(candidate):
    with pytest.raises(GovernanceBlock, match="no human approval token"):
        fcs_scale.promote_fcs_scale_regime(
            candidate,
            authority="EXPLICIT_CHAIRMAN_JUSTIFICATION",
            justification="because the Chairman said so",
        )


def test_the_calibration_token_does_not_authorise_this_parameter(candidate):
    with pytest.raises(GovernanceBlock, match="different namespace"):
        fcs_scale.promote_fcs_scale_regime(
            candidate,
            authority="EXPLICIT_CHAIRMAN_JUSTIFICATION",
            justification="j",
            approval_token="APPROVE_V3_CALIBRATION_PROMOTION::R9-SOMETHING",
        )


def test_mathematical_promotion_needs_holdout_evidence_and_a_mounted_dataset(candidate):
    with pytest.raises(GovernanceBlock, match="holdout split"):
        fcs_scale.promote_fcs_scale_regime(
            candidate,
            authority="GOVERNED_CALIBRATION_EVIDENCE",
            evidence={"split": "training", "dataset_sha256": "0" * 64},
        )
    with pytest.raises(GovernanceBlock, match="BLOCKED_ON_FCS_SCALE_EVIDENCE"):
        fcs_scale.promote_fcs_scale_regime(
            candidate,
            authority="GOVERNED_CALIBRATION_EVIDENCE",
            evidence={"split": "holdout"},
        )


def test_chairman_promotion_needs_a_written_justification(candidate):
    with pytest.raises(GovernanceBlock, match="written justification"):
        fcs_scale.promote_fcs_scale_regime(
            candidate, authority="EXPLICIT_CHAIRMAN_JUSTIFICATION", justification="   "
        )


def test_even_an_authorised_promotion_never_writes_the_canonical_config(candidate):
    result = fcs_scale.promote_fcs_scale_regime(
        candidate,
        authority="EXPLICIT_CHAIRMAN_JUSTIFICATION",
        justification="hypothetical, for the gate's own test",
        approval_token="APPROVE_V3_FCS_POINT_SCALE_PROMOTION::R9-HYPOTHETICAL",
    )
    assert result["writes_canonical_config"] is False
    assert fcs_scale.CANONICAL_FCS_UNIFIED_POINTS is None
    assert fcs.GOVERNED_FCS_POLICY.unified_points_equivalent is None


# --- exact evidence requirements ---------------------------------------------


def test_every_evidence_requirement_is_outstanding_and_names_a_closure_route():
    reqs = fcs_scale.FCS_POINT_SCALE_EVIDENCE_REQUIREMENTS
    assert len(reqs) >= 8
    assert all(r.satisfied is False for r in reqs)
    assert len(fcs_scale.unsatisfied_evidence_requirements()) == len(reqs)
    routes = {r.route for r in reqs}
    assert routes == {fcs_scale.ROUTE_ISSUED, fcs_scale.ROUTE_CALIBRATED, fcs_scale.ROUTE_EITHER}
    assert len({r.requirement_id for r in reqs}) == len(reqs)


def test_a_requirement_naming_an_unknown_route_is_refused():
    with pytest.raises(InputValidationError, match="unknown closure route"):
        fcs_scale.EvidenceRequirement(
            requirement_id="X", route="VIBES", requirement="", why_required=""
        )


def test_the_closed_routes_are_recorded_as_insufficient_rather_than_overlooked():
    text = " ".join(fcs_scale.INSUFFICIENT_EVIDENCE)
    assert "0.294" in text and "0.297514" in text
    assert "Inverting the POWER_CRUNCH Elo/Board transform" in text
    assert "category error" in text
    assert "affine, logistic or inverse transform" in text


# --- the blocker is retained --------------------------------------------------


def test_the_blocker_is_still_live_in_preflight(cfg):
    report = DynamicWeeklyMCV3(cfg).preflight()
    assert fcs.FCS_UNIFIED_SCALE_BLOCKER in report["execution_blockers"]
    status = report["r2_governance"]["fcs_point_scale_adapter"]
    assert status["blocker_retained"] is True
    assert status["adapter_interface_implemented"] is True
    assert status["adapter_constructed"] is False
    assert status["governed_for_execution"] is False
    assert status["canonical_unified_points"] is None


def test_the_lane_retired_no_blocker_and_opened_none():
    assert len(br.dispositions()) == 21
    entry = next(
        d for d in br.dispositions() if d.blocker_id == fcs.FCS_UNIFIED_SCALE_BLOCKER
    )
    assert entry.disposition == "MODEL_SCALE_ADAPTER_REQUIRED"
    assert entry.resolved is False
    assert entry.governance_group == br.GROUP_FCS_SCALE
    assert entry.ruling is None


def test_the_fcs_rating_policy_was_not_reopened():
    assert fcs.FCS_RATING_POLICY_RESOLVED is True
    assert fcs.FCS_FIXED_ELO == 1250.0
    assert fcs.require_governed_fcs_policy("FIXED_ELO_1250").fixed_elo == 1250.0
    status = fcs_scale.fcs_point_scale_status()
    assert status["rating_policy_resolved"] is True
    assert status["rating_policy_ruling"] == "R2-FCS-ELO-1250"


def test_the_engine_still_refuses_to_initialize_states(cfg, teams):
    with pytest.raises(GovernanceBlock, match="no unified preseason strength points"):
        DynamicWeeklyMCV3(cfg)._initialize_states(teams)


def test_the_canonical_config_carries_no_fcs_points_value():
    import json

    raw = json.loads(CONFIG.read_text(encoding="utf-8"))
    assert raw["fcs_translation_policy"] == "FIXED_ELO_1250"
    assert "fcs_unified_points_equivalent" not in raw
    assert "fcs_unified_points_equivalent" not in raw.get("calibration", {})


def test_the_calibration_harness_namespace_is_untouched_by_this_lane():
    from ncaaf_engine.simulation.dynamic_weekly_mc_v3 import calibration as cal

    assert cal.CALIBRATION_FIELDS == (
        "weekly_performance_residual_coefficient",
        "weekly_movement_cap_points",
        "recent_form_weights",
        "blowout_treatment",
        "game_sd_points",
        "sample_size_regularization",
    )
    assert not set(cal.CALIBRATION_FIELDS) & set(fcs_scale.FCS_SCALE_FIELDS)
    assert cal.load_candidate_regimes(CALIBRATION_REGIMES) == []


# --- the lane status artifact -------------------------------------------------

STATUS_F1 = ROOT / "reference/dynamic_weekly_mc_v3/V3_GOVERNANCE_STATUS_F1.json"


@pytest.fixture(scope="module")
def status_f1():
    import json

    return json.loads(STATUS_F1.read_text(encoding="utf-8"))


def test_the_lane_record_states_the_blocker_is_retained(status_f1):
    assert status_f1["outcome"] == "BLOCKER_RETAINED"
    assert status_f1["scope"] == "MODEL_SCALE_ADAPTER_LANE_ONLY"
    assert status_f1["base_sha"] == "eb5e3e7f546217809a94692bf0e502882b7d51c0"
    assert status_f1["head_sha"] is None
    assert status_f1["resolved_blockers"] == []
    assert status_f1["opened_blockers"] == []
    assert status_f1["chairman_rulings"] == []


# The lane record is epoch-specific and says so. F1 was frozen at eb5e3e7, where
# nine blockers were live; B1 independently retired the board-custody id off that
# same base. Both figures are true, of different epochs, and each is asserted
# against the epoch it names — never against the other.


def test_the_lane_record_states_its_own_nine_blocker_epoch(status_f1):
    assert status_f1["epoch_specific"] is True
    lane_base = sorted(br.F1_LANE_BASE_LIVE_BLOCKERS)
    assert status_f1["live_blockers_at_lane_base"] == lane_base
    assert status_f1["live_blocker_count_at_lane_base"] == len(lane_base) == 9
    assert status_f1["blocker_set_before"] == lane_base
    assert sorted(status_f1["live_blocker_classification_at_lane_base"]) == lane_base
    # The observation F1 recorded is preserved, board custody and all.
    assert "inputs.board_of_record_i_k" in status_f1["live_blockers_at_lane_base"]
    assert fcs.FCS_UNIFIED_SCALE_BLOCKER in status_f1["live_blockers_at_lane_base"]


def test_the_lane_record_matches_the_integrated_live_blocker_set(cfg, status_f1):
    report = DynamicWeeklyMCV3(cfg).preflight()
    live = sorted(str(b) for b in report["execution_blockers"])
    assert status_f1["integrated_live_blockers"] == live
    assert status_f1["integrated_live_blocker_count"] == len(live) == 8
    assert sorted(status_f1["integrated_live_blocker_classification"]) == live
    assert fcs.FCS_UNIFIED_SCALE_BLOCKER in live


def test_the_integrated_live_set_is_exactly_the_eight_expected(cfg):
    """Set equality in both directions against the named eight. No count check."""
    expected = {
        "calibration.weekly_performance_residual_coefficient",
        "calibration.weekly_movement_cap_points",
        "calibration.recent_form_weights",
        "calibration.blowout_treatment",
        "calibration.game_sd_points",
        "calibration.sample_size_regularization",
        "governance.GAME_SD_CALIBRATION_OPEN",
        "model_scale.FCS_ELO_1250_TO_V3_POINT_SCALE_ADAPTER",
    }
    live = {str(b) for b in DynamicWeeklyMCV3(cfg).preflight()["execution_blockers"]}
    assert not live - expected, "unexpected blocker live"
    assert not expected - live, "expected blocker missing"
    assert live == expected == set(br.F1_INTEGRATED_LIVE_BLOCKERS)


def test_this_lane_did_not_move_the_live_set_in_either_epoch(status_f1):
    delta = br.f1_integration_delta()
    assert delta["retired_by_f1"] == []
    assert delta["opened_by_f1"] == []
    # The one difference between the epochs is B1's, not F1's.
    assert delta["difference_from_lane_base"] == ["inputs.board_of_record_i_k"]
    assert delta["attributable_to_f1"] == []
    assert delta["attributable_to_b1"] == ["inputs.board_of_record_i_k"]
    assert delta["unexplained"] == []
    assert status_f1["blocker_epochs"]["attributable_to_f1"] == []
    assert status_f1["integration"]["blockers_retired_by_this_candidate"] == []
    assert status_f1["integration"]["blockers_opened_by_this_candidate"] == []


def test_the_f1_blocker_is_live_in_both_epochs(cfg, status_f1):
    delta = br.f1_integration_delta()
    assert delta["f1_blocker_live_at_lane_base"] is True
    assert delta["f1_blocker_live_when_integrated"] is True
    assert br.F1_MODEL_SCALE_BLOCKER == fcs.FCS_UNIFIED_SCALE_BLOCKER
    live = DynamicWeeklyMCV3(cfg).preflight()["execution_blockers"]
    assert fcs.FCS_UNIFIED_SCALE_BLOCKER in live
    assert status_f1["integration"]["f1_blocker_retained"] is True


def test_b1_custody_closure_survives_this_lane(cfg, status_f1):
    """F1 must not resurrect the blocker B1 closed on custody evidence."""
    delta = br.f1_integration_delta()
    assert delta["board_custody_live_at_lane_base"] is True
    assert delta["board_custody_live_when_integrated"] is False
    live = {str(b) for b in DynamicWeeklyMCV3(cfg).preflight()["execution_blockers"]}
    assert "inputs.board_of_record_i_k" not in live
    assert status_f1["integration"]["b1_custody_closure_preserved"] is True
    assert status_f1["integration"]["board_of_record_i_k_reopened_by_f1"] is False


def test_the_record_names_the_integration_base_and_its_own_provenance(status_f1):
    integration = status_f1["integration"]
    assert integration["original_f1_candidate_sha"] == (
        "0507dbe99e23b5b067e3af9305f4551194c2f83d"
    )
    assert integration["original_f1_base_sha"] == status_f1["base_sha"]
    assert integration["successor_head_sha"] is None
    assert integration["governance_authority_changed"] is False
    assert integration["new_ruling_issued"] is False
    assert integration["calibration_value_promoted"] is False
    assert integration["elo_to_points_transform_invented"] is False
    assert integration["conflict_resolved"]["ours_or_theirs_wholesale"] is False


def test_the_lane_record_ran_no_simulation_and_produced_no_probabilities(status_f1):
    assert status_f1["simulation_run"] is False
    assert status_f1["monte_carlo_10000_path_run"] is False
    assert status_f1["output_probabilities_generated"] is False
    assert status_f1["canonical_config"]["modified_by_f1"] is False
    assert status_f1["canonical_config"]["writes_canonical_config"] is False
    assert not (ROOT / "output").exists()


def test_the_lane_record_carries_the_reproductions_it_claims(status_f1):
    repro = status_f1["reproduced_transforms"]
    assert repro["do_they_compose"] is False
    elo = repro["elo_from_board_power_h"]
    assert elo["rows_reproduced"] == 121
    assert elo["reproduced_max_residual"] == fcs_scale.ELO_BOARD_REPRODUCED_MAX_RESIDUAL
    assert round(elo["reproduced_max_residual"], 4) == elo["recorded_max_residual"]
    points = repro["unified_neutral_field_points"]
    assert points["rows_reproduced"] == 121
    assert points["max_absolute_error"] == 0.0


def test_the_lane_record_names_the_manufactured_value_only_to_refuse_it(status_f1):
    closed = status_f1["closed_routes_refused_by_name"]
    assert closed["manufactured_value_provenance"] == "DERIVED_AND_REFUSED"
    assert closed["manufactured_board_equivalent_for_elo_1250"] == (
        fcs_scale.MANUFACTURED_BOARD_EQUIVALENT_FOR_ELO_1250
    )
    with pytest.raises(GovernanceBlock):
        fcs_scale.reject_manufactured_board_equivalent(
            closed["manufactured_board_equivalent_for_elo_1250"]
        )


def test_the_lane_record_carries_every_evidence_requirement(status_f1):
    ids = [r["requirement_id"] for r in status_f1["evidence_requirements"]]
    assert ids == [r.requirement_id for r in fcs_scale.FCS_POINT_SCALE_EVIDENCE_REQUIREMENTS]
    assert status_f1["unsatisfied_evidence_requirement_ids"] == ids
    assert all(r["satisfied"] is False for r in status_f1["evidence_requirements"])


def test_the_lane_record_records_the_governed_inputs_unchanged(status_f1):
    import hashlib

    for name, digest in status_f1["artifact_hashes"].items():
        path = INPUTS / name
        assert path.exists(), name
        assert hashlib.sha256(path.read_bytes()).hexdigest() == digest
    assert status_f1["governed_reference_inputs_unchanged"] is True


def test_the_lane_record_states_the_test_counts_it_was_measured_at(status_f1):
    counts = status_f1["test_counts"]
    # Arithmetic invariants only. Asserting a literal here would be self-referential:
    # the assertion itself is one of the tests being counted.
    assert counts["total"] == counts["predecessor_r4_total"] + counts["f1_lane"]
    assert counts["v3"] == counts["predecessor_r4_v3"] + counts["f1_lane"]


def test_the_lane_record_states_the_counts_of_the_integrated_tree(status_f1):
    """The historical block above is F1's own; this one is the integrated tree's."""
    counts = status_f1["integrated_test_counts"]
    assert counts["total"] == (
        counts["predecessor_integrated_base_total"] + counts["f1_lane"]
    )
    assert counts["v3"] == counts["predecessor_integrated_base_v3"] + counts["f1_lane"]
    assert counts["f1_lane"] != status_f1["test_counts"]["f1_lane"]


def test_the_lane_record_agrees_with_the_module_coverage(status_f1, coverage):
    recorded = status_f1["coverage"]
    assert recorded["schedule_only_fcs_identity_count"] == coverage["fcs_identity_count"] == 13
    assert recorded["fbs_v_fcs_game_count"] == coverage["fbs_v_fcs_game_count"] == 15
    assert recorded["fcs_hosted_game_ids"] == coverage["fcs_home_game_ids"]
    assert recorded["identities_resolvable_on_points_axis"] == 0
    assert recorded["games_simulatable"] == 0
