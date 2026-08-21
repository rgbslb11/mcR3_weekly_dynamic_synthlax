from pathlib import Path

import pytest
from openpyxl import load_workbook

from ncaaf_engine.simulation.dynamic_weekly_mc_v3.config import V3Config
from ncaaf_engine.simulation.dynamic_weekly_mc_v3.inputs import schedule_content_hash, sha256_file
from ncaaf_engine.simulation.dynamic_weekly_mc_v3.provenance import (
    BINARY_MISMATCH,
    CERTIFIED_GAMES_CONTENT_SHA256,
    CONTENT_MISMATCH,
    GOVERNED_SCHEDULE_BINARY_SHA256,
    SUPERSEDED_V4_BINARY_SHA256,
    ScheduleProvenance,
    reconcile_schedule_provenance,
)

ROOT = Path(__file__).resolve().parents[2]
CONFIG = ROOT / "config/dynamic_weekly_mc_v3/v3_experimental.json"


def _schedule() -> Path:
    return V3Config.from_json(CONFIG).inputs.schedule_xlsx


def test_certified_games_content_hash_is_reproducible():
    """The canonical serialization must reproduce the certification exactly."""
    assert schedule_content_hash(_schedule()) == CERTIFIED_GAMES_CONTENT_SHA256


def test_ragged_rows_are_padded_to_header_width():
    """735 of 743 data rows omit the trailing cell; they must serialize as blank."""
    wb = load_workbook(_schedule(), read_only=True, data_only=True)
    rows = [list(r) for r in wb["Games"].iter_rows(values_only=True) if any(v is not None for v in r)]
    width = len(rows[0])
    short = [r for r in rows[1:] if len(r) < width]
    assert width == 13
    assert len(short) == 735, "row-width profile changed; the padding rule needs re-verification"


def test_binary_hash_does_not_match_registered_v5_artifact():
    """Recorded so a later fix cannot silently change custody state."""
    actual = sha256_file(_schedule())
    assert actual != GOVERNED_SCHEDULE_BINARY_SHA256
    assert actual != SUPERSEDED_V4_BINARY_SHA256, "mounted file is the superseded v4 artifact"


def test_reconciliation_classifies_content_verified_binary_mismatch():
    prov = reconcile_schedule_provenance(_schedule())
    assert prov.content_verified is True
    assert prov.binary_verified is False
    assert prov.classification() == "CONTENT_VERIFIED_BINARY_SOURCE_COPY_MISMATCH"
    assert prov.anomalies() == [BINARY_MISMATCH]


def test_mutated_games_content_fails_the_certification_gate(tmp_path):
    """A fixture edit must break the content hash even though the file re-saves cleanly."""
    mutant = tmp_path / "mutant.xlsx"
    mutant.write_bytes(_schedule().read_bytes())
    wb = load_workbook(mutant)
    wb["Games"]["E2"] = "AUB"
    wb.save(mutant)
    assert schedule_content_hash(mutant) != CERTIFIED_GAMES_CONTENT_SHA256
    prov = reconcile_schedule_provenance(mutant)
    assert prov.classification() == "CANONICAL_CONTENT_MISMATCH"
    assert CONTENT_MISMATCH in prov.anomalies()


def test_resaved_workbook_keeps_content_and_loses_binary(tmp_path):
    """The asymmetry the classifier depends on: content survives a re-save, binary does not."""
    resaved = tmp_path / "resaved.xlsx"
    resaved.write_bytes(_schedule().read_bytes())
    wb = load_workbook(resaved)
    wb.save(resaved)
    assert schedule_content_hash(resaved) == CERTIFIED_GAMES_CONTENT_SHA256
    assert sha256_file(resaved) != sha256_file(_schedule())


@pytest.mark.parametrize(
    "content_ok, binary_hash, expected",
    [
        (True, GOVERNED_SCHEDULE_BINARY_SHA256, "FULLY_VERIFIED"),
        (False, GOVERNED_SCHEDULE_BINARY_SHA256, "CANONICAL_CONTENT_MISMATCH"),
        (True, SUPERSEDED_V4_BINARY_SHA256, "SUPERSEDED_V4_ARTIFACT_MOUNTED"),
        (True, "0" * 64, "CONTENT_VERIFIED_BINARY_SOURCE_COPY_MISMATCH"),
    ],
)
def test_classification_matrix(content_ok, binary_hash, expected):
    prov = ScheduleProvenance(
        content_sha256_reproduced=CERTIFIED_GAMES_CONTENT_SHA256 if content_ok else "f" * 64,
        content_sha256_certified=CERTIFIED_GAMES_CONTENT_SHA256,
        binary_sha256_actual=binary_hash,
        binary_sha256_governed=GOVERNED_SCHEDULE_BINARY_SHA256,
    )
    assert prov.classification() == expected
