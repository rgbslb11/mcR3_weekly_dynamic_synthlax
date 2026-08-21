from pathlib import Path

import pytest

from ncaaf_engine.simulation.dynamic_weekly_mc_v3.config import DEFAULT_PRIOR_DECAY, V3Config
from ncaaf_engine.simulation.dynamic_weekly_mc_v3.engine import DynamicWeeklyMCV3
from ncaaf_engine.simulation.dynamic_weekly_mc_v3.errors import GovernanceBlock, InputValidationError
from ncaaf_engine.simulation.dynamic_weekly_mc_v3.inputs import load_canonical_team_index, load_schedule, validate_schedule

ROOT = Path(__file__).resolve().parents[2]
CONFIG = ROOT / "config/dynamic_weekly_mc_v3/v3_experimental.json"


def test_governed_architecture_and_blockers_are_explicit():
    cfg = V3Config.from_json(CONFIG)
    cfg.validate_architecture()
    assert cfg.paths == 10_000
    assert cfg.weeks == tuple(range(1, 17))
    assert cfg.prior_decay == DEFAULT_PRIOR_DECAY
    assert cfg.audit_only_after_week(1) is True
    assert cfg.audit_only_after_week(2) is False
    blockers = set(cfg.execution_blockers())
    assert "calibration.weekly_performance_residual_coefficient" in blockers
    assert "calibration.weekly_movement_cap_points" in blockers
    assert "calibration.recent_form_weights" in blockers
    assert "calibration.blowout_treatment" in blockers
    assert "calibration.game_sd_points" in blockers
    assert "calibration.sample_size_regularization" in blockers
    assert "hfa_baseline_points" in blockers
    assert "fcs_translation_policy" in blockers
    assert "committee_tiebreak_strength_source" in blockers
    assert "inputs.aac_divisions_csv" in blockers
    with pytest.raises(GovernanceBlock):
        cfg.require_executable()


def test_real_inputs_pass_structural_preflight():
    cfg = V3Config.from_json(CONFIG)
    report = DynamicWeeklyMCV3(cfg).preflight()
    assert report["status"] == "STRUCTURAL_PREFLIGHT_PASS"
    assert report["team_entities"] == 134
    assert report["fbs_members"] == 121
    assert report["schedule_only_fcs"] == 13
    assert report["schedule"]["games_total"] == 743
    assert report["schedule"]["regular_games"] == 736
    assert report["schedule"]["ccg_templates"] == 7
    assert report["schedule"]["schedule_games_sha256_certified"] == "bd8089f70f6d483a75564e33438272c22daade8e53619fb21a915778975ff221"
    # The Games-sheet content hash now reproduces the certification exactly; the
    # binary artifact still is not the registered v5 upload.
    assert report["schedule"]["schedule_games_sha256_reproduced"] == report["schedule"]["schedule_games_sha256_certified"]
    assert "SCHEDULE_GAMES_HASH_REPRODUCTION_MISMATCH" not in report["schedule"]["provenance_anomalies"]
    assert "SCHEDULE_BINARY_HASH_MISMATCH_VS_MODEL_PARAMETERS_V2_5" in report["schedule"]["provenance_anomalies"]


def test_canonical_team_index_is_exactly_134():
    cfg = V3Config.from_json(CONFIG)
    rows = load_canonical_team_index(cfg.inputs.canonical_master_md)
    assert len(rows) == 134
    assert sum(r["entity_scope"] == "FBS_MEMBER" for r in rows.values()) == 121
    assert sum(r["entity_scope"] == "SCHEDULE_ONLY_FCS" for r in rows.values()) == 13


def test_schedule_mutation_fails_hash_gate(tmp_path):
    cfg = V3Config.from_json(CONFIG)
    from openpyxl import load_workbook
    mutant = tmp_path / "mutant.xlsx"
    mutant.write_bytes(cfg.inputs.schedule_xlsx.read_bytes())
    wb = load_workbook(mutant)
    ws = wb["Games"]
    ws["E2"] = "AUB"
    wb.save(mutant)
    from ncaaf_engine.simulation.dynamic_weekly_mc_v3.inputs import schedule_content_hash
    assert schedule_content_hash(mutant) != schedule_content_hash(cfg.inputs.schedule_xlsx)


def test_powershell_launcher_targets_v3_cli():
    launcher = ROOT / "scripts/Invoke-SythalaxDynamicWeeklyMCV3.ps1"
    text = launcher.read_text(encoding="utf-8")
    assert "ncaaf_engine.simulation.dynamic_weekly_mc_v3.cli" in text
    assert "-Command validate" not in text  # command is parameterized, not hard-wired
    assert "show-blockers" in text and "run" in text


def test_blocked_run_emits_no_final_looking_outputs(tmp_path):
    from ncaaf_engine.simulation.dynamic_weekly_mc_v3.cli import main

    rc = main([
        "run",
        "--config", str(CONFIG),
        "--output-root", str(tmp_path),
    ])
    assert rc == 2
    files = [p.relative_to(tmp_path).as_posix() for p in tmp_path.rglob("*") if p.is_file()]
    assert len(files) == 1
    assert files[0].endswith("/BLOCKED.json")
    forbidden = ("team_summary", "conference_summary", "cfp", "probability", "parquet")
    assert not any(any(token in name.lower() for token in forbidden) for name in files)
