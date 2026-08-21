"""FCS Elo 1250 -> V3 unified neutral-point scale adapter.

The FCS *rating policy* is settled at Elo 1250 by ruling R2-FCS-ELO-1250 and is
not reopened anywhere in this module. What is under test is the separate
model-scale question: whether V3 may obtain a unified neutral-point value for
the 13 schedule-only FCS entities, and by what route.

The answer these tests pin down is *no route currently exists*, and every
plausible substitute is refused by name. A test here that starts passing because
someone supplied a number has not been fixed, it has been defeated.
"""

from __future__ import annotations

import inspect
from pathlib import Path

import pytest
from openpyxl import load_workbook

from ncaaf_engine.simulation.dynamic_weekly_mc_v3 import engine as engine_module
from ncaaf_engine.simulation.dynamic_weekly_mc_v3 import fcs
from ncaaf_engine.simulation.dynamic_weekly_mc_v3.config import V3Config
from ncaaf_engine.simulation.dynamic_weekly_mc_v3.engine import DynamicWeeklyMCV3
from ncaaf_engine.simulation.dynamic_weekly_mc_v3.errors import GovernanceBlock
from ncaaf_engine.simulation.dynamic_weekly_mc_v3.inputs import load_schedule, load_teams

ROOT = Path(__file__).resolve().parents[2]
CONFIG = ROOT / "config" / "dynamic_weekly_mc_v3" / "v3_experimental.json"
INPUTS = ROOT / "reference" / "dynamic_weekly_mc_v3" / "inputs"

FCS_ENTITIES = {
    "ARST", "CHAR", "CP", "DUQ", "EMU", "IDHO", "SAC",
    "SUU", "TOL", "ULL", "ULM", "WKU", "WMU",
}


@pytest.fixture(autouse=True)
def _no_adapter_leaks():
    """No test may leave an adapter installed for the next one."""
    fcs.clear_fcs_scale_adapter()
    yield
    fcs.clear_fcs_scale_adapter()


@pytest.fixture(scope="module")
def teams():
    cfg = V3Config.from_json(CONFIG)
    return load_teams(cfg.inputs.canonical_master_md, cfg.inputs.unified_preseason_ratings_xlsx)


def _code_of(func) -> str:
    """Source with comments stripped, so a comment *about* a fallback cannot
    trip a check that is looking for the fallback itself."""
    lines = []
    for line in inspect.getsource(func).splitlines():
        stripped = line.split("#", 1)[0]
        if stripped.strip():
            lines.append(stripped)
    return "\n".join(lines)


def _governed_provenance(**overrides):
    base = dict(
        artifact="R5-FCS-SCALE (hypothetical)",
        locator="Rulings_Register!A1",
        authority="DIRECT_CHAIRMAN_AUTHORITY",
        issued=True,
        statement="FCS Elo 1250 corresponds to N unified neutral-field points.",
    )
    base.update(overrides)
    return fcs.FcsScaleProvenance(**base)


def _adapter(points=-9.0, **overrides):
    kwargs = dict(
        unified_points=points,
        provenance=_governed_provenance(),
        derivation="DIRECT_CHAIRMAN_AUTHORITY",
    )
    kwargs.update(overrides)
    return fcs.FcsScaleAdapter(**kwargs)


# --- Census: the problem is the size the docstring claims ---------------------


def test_the_schedule_only_fcs_universe_is_exactly_thirteen_entities(teams):
    found = {sid for sid, t in teams.items() if t.entity_scope == "SCHEDULE_ONLY_FCS"}
    assert found == FCS_ENTITIES
    assert len(found) == 13
    assert all(teams[sid].preseason_strength_points is None for sid in found)


def test_exactly_fifteen_regular_season_games_involve_an_fcs_entity():
    schedule = load_schedule(INPUTS / "2026_FBS_Schedule_LOCKED_v5.xlsx")
    regular = [g for g in schedule if g.game_type == "REG"]
    involving = [
        g for g in regular
        if g.home_team in FCS_ENTITIES or g.away_team in FCS_ENTITIES
    ]
    assert len(involving) == 15
    assert sorted({g.week for g in involving}) == [2, 3, 4, 5, 12]
    # The fcs_game flag and the canonical entity set must agree exactly, so the
    # census cannot be satisfied by trusting either one alone.
    assert {g.game_id for g in involving} == {g.game_id for g in regular if g.fcs_game}
    # Every one of the 13 is actually scheduled; none is a dormant row.
    assert {
        t for g in involving for t in (g.home_team, g.away_team) if t in FCS_ENTITIES
    } == FCS_ENTITIES


# --- The target axis, and why FCS is not on it -------------------------------


def test_the_unified_point_axis_is_a_z_score_of_the_fbs_population_only():
    """Reproduce the axis from the governed workbook rather than asserting it."""
    wb = load_workbook(
        INPUTS / "2026_CFB_2026_Preseason_Unified_Power_Ratings.xlsx",
        read_only=True, data_only=True,
    )
    ws = wb["Master Ratings"]
    rows = [r for r in ws.iter_rows(min_row=5, values_only=True) if r[4] is not None]

    assert len(rows) == fcs.UNIFIED_Z_POPULATION_SIZE == 121
    # Points are exactly 14 x Unified Master Z.
    assert max(
        abs(float(r[72]) - fcs.UNIFIED_NEUTRAL_POINTS_PER_SD * float(r[70])) for r in rows
    ) < 1e-9
    # Unified Master Z is the unweighted mean of the four family Z scores.
    assert max(
        abs(float(r[70]) - (float(r[64]) + float(r[65]) + float(r[69]) + float(r[68])) / 4)
        for r in rows
    ) < 1e-9
    # Board Family Z is itself the mean of Board I-H and Board J-B, so Board I-H
    # carries 0.125 of the axis and no more.
    assert max(
        abs(float(r[68]) - (float(r[66]) + float(r[67])) / 2) for r in rows
    ) < 1e-9
    assert fcs.BOARD_IH_SHARE_OF_UNIFIED_Z == 0.25 * 0.5

    # Not one FCS entity appears in the rated population.
    assert not FCS_ENTITIES & {str(r[4]) for r in rows}


def test_even_the_forbidden_board_route_could_not_reconstruct_a_point_value():
    """The refusal does not rest on the ruling alone.

    Board I-H supplies 0.125 of Unified Master Z. The other 0.875 — TrueSkill,
    Litkenhous, Pure Baxter and Board J-B — has no governed value for any FCS
    entity, so even if the Board route were permitted it would still be short of
    the inputs the axis is made of.
    """
    assert fcs.UNGOVERNED_SHARE_OF_AXIS_FOR_FCS == pytest.approx(0.875)
    assert fcs.BOARD_IH_SHARE_OF_UNIFIED_Z + fcs.UNGOVERNED_SHARE_OF_AXIS_FOR_FCS == 1.0

    wb = load_workbook(
        INPUTS / "POWER_CRUNCH_2026_Team_Reconciled_Master_134_v2_2.xlsx",
        read_only=True, data_only=True,
    )
    rows = wb["Reconciled Master"].iter_rows(values_only=True)
    header = [str(x) if x is not None else "" for x in next(rows)]
    idx = {name: i for i, name in enumerate(header)}
    fcs_rows = [r for r in rows if r[idx["entity_scope"]] == "SCHEDULE_ONLY_FCS"]

    assert len(fcs_rows) == 13
    # The governed source records sentinels, not numbers.
    assert {r[idx["board_power_H"]] for r in fcs_rows} == {fcs.FCS_BOARD_POWER_H_SENTINEL}
    assert {
        r[idx["home_field_advantage_modifier"]] for r in fcs_rows
    } == {fcs.FCS_HFA_MODIFIER_SENTINEL}


# --- Required refusals -------------------------------------------------------


def test_direct_1250_as_points_is_refused():
    with pytest.raises(GovernanceBlock, match="different axis"):
        fcs.reject_elo_as_points(fcs.FCS_FIXED_ELO)
    with pytest.raises(GovernanceBlock, match=fcs.FCS_UNIFIED_SCALE_BLOCKER):
        fcs.register_fcs_scale_adapter(
            _adapter(points=1250.0),
            approval_token="APPROVE_V3_FCS_SCALE_ADAPTER::R5-HYPOTHETICAL",
        )
    assert fcs.active_fcs_scale_adapter() is None


def test_the_superseded_operator_composite_elo_is_also_refused_as_points():
    assert fcs.SUPERSEDED_OPERATOR_COMPOSITE_ELO == 1397.51
    with pytest.raises(GovernanceBlock, match="different axis"):
        fcs.reject_elo_as_points(1397.51)


@pytest.mark.parametrize("value", [0.294, 0.297, 0.297514])
def test_the_enumerated_board_equivalents_are_refused(value):
    with pytest.raises(GovernanceBlock, match="conversion rule"):
        fcs.reject_board_derived_conversion(value)
    with pytest.raises(GovernanceBlock, match="conversion rule"):
        fcs.register_fcs_scale_adapter(
            _adapter(points=value),
            approval_token="APPROVE_V3_FCS_SCALE_ADAPTER::R5-HYPOTHETICAL",
        )


def test_the_board_inverse_of_the_governed_elo_is_refused_even_though_no_list_names_it():
    """The gap a three-value blacklist left open.

    0.294 / 0.297 / 0.297514 are all inversions of the *superseded* 1397.51.
    Inverting the *governed* 1250 yields roughly 0.150002, which appears on no
    list, and a value-only refusal would have admitted it.
    """
    smuggled = fcs.board_inverse_of(fcs.FCS_FIXED_ELO)
    assert smuggled == pytest.approx(0.150002, abs=1e-6)
    assert round(smuggled, 6) not in {round(v, 6) for v in fcs.FORBIDDEN_BOARD_EQUIVALENTS}
    with pytest.raises(GovernanceBlock, match="inverted Elo/Board transform"):
        fcs.reject_board_derived_conversion(smuggled)


def test_inverting_the_board_transform_is_refused_as_a_method():
    with pytest.raises(GovernanceBlock, match="forbidden"):
        fcs.invert_board_transform(fcs.FCS_FIXED_ELO)
    for derivation in ("BOARD_IH_INVERSE", "BOARD_EQUIVALENT", "V2_1_BRIDGE_REPLAY"):
        with pytest.raises(GovernanceBlock, match="refused as a"):
            fcs.register_fcs_scale_adapter(
                _adapter(derivation=derivation),
                approval_token="APPROVE_V3_FCS_SCALE_ADAPTER::R5-HYPOTHETICAL",
            )
    assert fcs.active_fcs_scale_adapter() is None


def test_an_ad_hoc_or_test_fitted_coefficient_is_refused():
    for derivation in ("AD_HOC", "TEST_FITTING", "ELO_AS_POINTS"):
        with pytest.raises(GovernanceBlock, match="refused as a"):
            fcs.register_fcs_scale_adapter(
                _adapter(derivation=derivation),
                approval_token="APPROVE_V3_FCS_SCALE_ADAPTER::R5-HYPOTHETICAL",
            )


def test_an_unrecognised_derivation_is_refused_rather_than_waved_through():
    with pytest.raises(GovernanceBlock, match="unknown FCS scale derivation"):
        fcs.register_fcs_scale_adapter(
            _adapter(derivation="SEEMS_REASONABLE"),
            approval_token="APPROVE_V3_FCS_SCALE_ADAPTER::R5-HYPOTHETICAL",
        )


@pytest.mark.parametrize("field", ["artifact", "locator", "statement"])
def test_missing_provenance_is_refused(field):
    with pytest.raises(GovernanceBlock, match="missing"):
        fcs.register_fcs_scale_adapter(
            _adapter(provenance=_governed_provenance(**{field: "  "})),
            approval_token="APPROVE_V3_FCS_SCALE_ADAPTER::R5-HYPOTHETICAL",
        )
    assert fcs.active_fcs_scale_adapter() is None


def test_a_recorded_but_unissued_provenance_is_refused():
    """'Recorded not issued' is the distinction the whole blocker turns on."""
    with pytest.raises(GovernanceBlock, match="recorded, "):
        fcs.register_fcs_scale_adapter(
            _adapter(provenance=_governed_provenance(issued=False)),
            approval_token="APPROVE_V3_FCS_SCALE_ADAPTER::R5-HYPOTHETICAL",
        )


def test_an_unrecognised_provenance_authority_is_refused():
    with pytest.raises(GovernanceBlock, match="unknown FCS scale authority"):
        _governed_provenance(authority="OPERATOR_ASSIGNED_COMPOSITE").validate()


def test_a_margin_calibrated_adapter_must_name_its_experiment():
    with pytest.raises(GovernanceBlock, match="must name the calibration experiment"):
        fcs.register_fcs_scale_adapter(
            _adapter(
                derivation="MARGIN_CALIBRATED_HOLDOUT",
                provenance=_governed_provenance(authority="MARGIN_CALIBRATED_HOLDOUT"),
                calibration_id=None,
            ),
            approval_token="APPROVE_V3_FCS_SCALE_ADAPTER::R5-HYPOTHETICAL",
        )


def test_a_structurally_valid_adapter_still_needs_an_approval_token():
    with pytest.raises(GovernanceBlock, match="no human approval token"):
        fcs.register_fcs_scale_adapter(_adapter())
    with pytest.raises(GovernanceBlock, match="malformed"):
        fcs.register_fcs_scale_adapter(_adapter(), approval_token="yes please")
    assert fcs.active_fcs_scale_adapter() is None


# --- Missing calibrated adapter: the live state ------------------------------


def test_no_adapter_is_installed_and_the_accessor_fails_closed():
    assert fcs.active_fcs_scale_adapter() is None
    assert fcs.fcs_unified_scale_governed() is False
    assert fcs.GOVERNED_FCS_POLICY.unified_points_equivalent is None
    with pytest.raises(GovernanceBlock, match=fcs.FCS_UNIFIED_SCALE_BLOCKER):
        fcs.require_fcs_unified_points()


def test_the_blocker_is_live_and_is_not_retired():
    report = DynamicWeeklyMCV3(V3Config.from_json(CONFIG)).preflight()
    assert fcs.FCS_UNIFIED_SCALE_BLOCKER in report["execution_blockers"]
    assert report["r2_governance"]["fcs"]["scale_adapter_installed"] is False
    assert report["r2_governance"]["fcs"]["unified_points_equivalent"] is None
    # The settled policy is reported alongside it and is untouched.
    assert report["r2_governance"]["fcs"]["fixed_elo"] == 1250.0
    assert report["r2_governance"]["fcs"]["rating_policy_resolved"] is True


def test_initialize_states_refuses_and_names_every_unrated_entity(teams):
    with pytest.raises(GovernanceBlock, match="no unified preseason strength points") as exc:
        DynamicWeeklyMCV3(V3Config.from_json(CONFIG))._initialize_states(teams)
    message = str(exc.value)
    assert "13 entities" in message
    assert all(sid in message for sid in FCS_ENTITIES)
    # It names the blocker, so the failure is traceable to the lane that owns it.
    assert fcs.FCS_UNIFIED_SCALE_BLOCKER in message


# --- Hidden fallbacks --------------------------------------------------------


def test_the_engine_has_no_silent_default_for_an_unrated_entity(teams):
    """There must be exactly one route to an FCS point value, and it must refuse.

    Previously the refusal lived inline in ``_initialize_states`` and the module
    accessor had no production caller at all, so the two could drift apart.
    """
    source = _code_of(engine_module.DynamicWeeklyMCV3._initialize_states)
    assert "require_fcs_unified_points" in source
    assert "or 1.0" not in source
    assert "0.0" not in source

    # And it genuinely refuses rather than defaulting, for every one of the 13.
    for sid in FCS_ENTITIES:
        assert teams[sid].preseason_strength_points is None


def test_the_fcs_home_field_modifier_is_not_silently_defaulted_to_one():
    """Three scheduled games put an FCS entity at a HOME venue.

    POWER_CRUNCH records their home_field_advantage_modifier as UNRESOLVED, so
    ``float(modifier or 1.0)`` invented a governed quantity for real games.
    """
    schedule = load_schedule(INPUTS / "2026_FBS_Schedule_LOCKED_v5.xlsx")
    fcs_at_home = [
        g for g in schedule
        if g.game_type == "REG" and g.home_team in FCS_ENTITIES and g.venue == "HOME"
    ]
    assert {g.game_id for g in fcs_at_home} == {"G0019", "G0213", "G0224"}

    with pytest.raises(GovernanceBlock, match="no governed home-field modifier"):
        fcs.require_fcs_hfa_modifier(None, "TOL")
    assert "or 1.0" not in _code_of(
        engine_module.DynamicWeeklyMCV3.simulate_preselection_regular_season_path
    )
    # A real modifier still passes through, including a legitimate zero, which
    # the old ``or 1.0`` would have rewritten.
    assert fcs.require_fcs_hfa_modifier(1.0, "ALA") == 1.0
    assert fcs.require_fcs_hfa_modifier(0.0, "ALA") == 0.0


def test_no_module_constant_quietly_holds_an_fcs_point_value():
    forbidden = {round(v, 6) for v in fcs.FORBIDDEN_BOARD_EQUIVALENTS}
    forbidden |= {round(fcs.board_inverse_of(e), 6) for e in fcs.BOARD_INVERTED_ELOS}
    for name, value in vars(fcs).items():
        if name.startswith("_") or not isinstance(value, float):
            continue
        # Elo-layer constants are allowed to exist; point-layer ones are not.
        if name in {
            "FCS_FIXED_ELO",
            "SUPERSEDED_OPERATOR_COMPOSITE_ELO",
            "POWER_CRUNCH_ELO_BOARD_SLOPE",
            "POWER_CRUNCH_ELO_BOARD_INTERCEPT",
            "UNIFIED_NEUTRAL_POINTS_PER_SD",
            "BOARD_IH_SHARE_OF_UNIFIED_Z",
            "UNGOVERNED_SHARE_OF_AXIS_FOR_FCS",
        }:
            continue
        assert round(value, 6) not in forbidden, f"{name} holds a forbidden conversion value"


# --- The interface is usable, not merely a wall ------------------------------


def test_a_fully_governed_adapter_installs_and_then_satisfies_the_engine():
    """Proof the gate is a door.

    The value below is a placeholder attached to a hypothetical ruling id. It is
    not proposed, not promoted, and is removed by the fixture the moment this
    test ends. Its only purpose is to show that a genuinely issued adapter has a
    route in, so the refusals above are a governance gate rather than a dead end.
    """
    adapter = fcs.register_fcs_scale_adapter(
        _adapter(points=-9.0),
        approval_token="APPROVE_V3_FCS_SCALE_ADAPTER::R5-HYPOTHETICAL",
    )
    assert fcs.active_fcs_scale_adapter() is adapter
    assert fcs.fcs_unified_scale_governed() is True
    assert fcs.require_fcs_unified_points() == -9.0
    # The settled Elo policy is untouched by installing a scale adapter.
    assert fcs.GOVERNED_FCS_POLICY.fixed_elo == 1250.0
    assert fcs.GOVERNED_FCS_POLICY.unified_points_equivalent is None


def test_clearing_the_adapter_restores_the_fail_closed_state():
    fcs.register_fcs_scale_adapter(
        _adapter(), approval_token="APPROVE_V3_FCS_SCALE_ADAPTER::R5-HYPOTHETICAL"
    )
    fcs.clear_fcs_scale_adapter()
    assert fcs.fcs_unified_scale_governed() is False
    with pytest.raises(GovernanceBlock):
        fcs.require_fcs_unified_points()


# --- Evidence report ---------------------------------------------------------


def test_the_evidence_report_promotes_no_number_and_shows_the_search_was_exhaustive():
    report = fcs.fcs_scale_evidence_report()
    assert report["numeric_mapping_promoted"] is False
    assert report["adapter_installed"] is False
    assert report["disposition"] == "BLOCKED_ON_NUMERICAL_SCALE_CALIBRATION"
    assert report["fcs_entity_count"] == 13
    assert report["fcs_regular_season_games"] == 15
    assert report["rating_policy_resolved"] is True
    assert report["fixed_elo"] == 1250.0

    routes = {r["route"] for r in report["candidate_routes"]}
    assert routes == {
        "ELO_AS_POINTS",
        "BOARD_EQUIVALENT",
        "BOARD_IH_INVERSE",
        "V2_1_BRIDGE_REPLAY",
        "ISSUED_GOVERNED_REGISTER",
        "MARGIN_CALIBRATED_HOLDOUT",
    }
    # Every route is refused with a stated reason and carries no number.
    for route in report["candidate_routes"]:
        assert route["numeric_result"] is None
        assert route["refused_because"].strip()


def test_the_axis_scale_itself_is_recorded_as_pending_calibration():
    """The 14 points/SD scale is provisional in its own source.

    Any FCS point value would be calibrated against an axis whose scale is
    itself open, which is why the adapter is a calibration item and not merely a
    missing constant.
    """
    report = fcs.fcs_scale_evidence_report()
    assert report["target_axis"]["scale_status"] == "initial scale pending margin calibration"

    wb = load_workbook(
        INPUTS / "2026_CFB_2026_Preseason_Unified_Power_Ratings.xlsx",
        read_only=True, data_only=True,
    )
    dictionary = {
        str(r[0]): str(r[1])
        for r in wb["Data Dictionary"].iter_rows(values_only=True)
        if r[0] is not None
    }
    assert "pending margin calibration" in dictionary["Unified Neutral-Field Points"]
