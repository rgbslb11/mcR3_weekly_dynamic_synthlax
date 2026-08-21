"""Historical calibration data plane for Dynamic Weekly MC V3.

The calibration harness in :mod:`.calibration` can score a candidate regime the
moment a historical observation set exists. This module is the layer underneath
it: it establishes *whether* such a set exists, loads it deterministically when
it does, and refuses to manufacture one when it does not.

The distinction this module exists to hold is between three things that a flat
reading of the mounted artifacts blurs together:

``fixtures``
    Schedule v5 lists 743 games of the 2026 season. A fixture names who plays
    whom; it carries no score. An unplayed season cannot supply an observation.

``derived ratings``
    The unified preseason workbook carries ``2025 Baxter Ridge Rating`` for all
    121 FBS teams, fitted over a recorded per-team game count. Those are season
    aggregates *computed from* 2025 games. The games themselves are not mounted,
    and a fitted rating cannot be un-fitted back into the observations that
    produced it.

``simulated output``
    The V2.1 static control is 10,000 simulated paths over the 2026 schedule,
    and the Model Parameters ``CALIBRATION`` sheet records achieved metrics from
    a prior *engine run*. Simulated margins are engine behaviour, not evidence
    about the world, and calibrating an engine against its own output is
    circular.

None of the three is a historical observation, so this module's honest answer on
the mounted evidence is :data:`MISSING_CALIBRATION_EVIDENCE` rather than a
dataset. The loaders, checks and split logic are complete and tested against
synthetic fixtures so that mounting a real set is a data step and not a build
step — but plumbing that works is not evidence that exists, and nothing here
retires a calibration blocker.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
from dataclasses import dataclass, replace
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Sequence

from .calibration import (
    BLOCKED_ON_CALIBRATION_DATA,
    CALIBRATION_FIELDS,
    CALIBRATION_OBSERVATION_COLUMNS,
    DATA_SPLITS,
    register_dataset,
)
from .errors import GovernanceBlock, InputValidationError

#: Emitted when the mounted evidence cannot support a calibration observation set.
MISSING_CALIBRATION_EVIDENCE = "MISSING_CALIBRATION_EVIDENCE"

#: A synthetic row is stamped in ``source_provenance`` rather than in a new
#: column: the observation allowlist is governed, and quietly widening it to
#: carry a test flag is exactly the "admit a signal by silence" move it forbids.
TEST_FIXTURE_STAMP = "TEST_FIXTURE::"


# --- mission schema -> governed column names ---------------------------------
#
# The lane schema and the governed allowlist name the same fields differently.
# Mapping them explicitly is what keeps the allowlist closed: every required
# field already has a governed column, so nothing has to be added by fiat.

MISSION_FIELD_TO_GOVERNED_COLUMN: dict[str, str] = {
    "game_id": "game_id",
    "season": "season",
    "week": "week",
    "event_time": "event_time",
    "team_id": "team",
    "opponent_id": "opponent",
    "site": "venue",
    "pregame_team_strength": "pregame_team_rating",
    "pregame_opponent_strength": "pregame_opponent_rating",
    "expected_margin": "expected_margin",
    "actual_margin": "actual_margin",
    "result": "game_result",
    "prior_rating_state": "prior_rating_state",
    "subsequent_outcomes": "subsequent_outcomes",
    "source_artifact_hash": "source_provenance",
    "observed_at": "observed_at",
    "recorded_at": "recorded_at",
    "model_version": "model_version",
    "config_version": "configuration_version",
}

#: Every column the data plane requires. Stricter than the harness minimum in
#: :mod:`.calibration`, which stays exactly as governed.
REQUIRED_DATA_PLANE_COLUMNS: tuple[str, ...] = (
    "game_id",
    "season",
    "week",
    "event_time",
    "team",
    "opponent",
    "venue",
    "pregame_team_rating",
    "pregame_opponent_rating",
    "expected_margin",
    "actual_margin",
    "game_result",
    "prior_rating_state",
    "source_provenance",
    "observed_at",
    "recorded_at",
    "model_version",
    "configuration_version",
)

# Import-time invariant: the lane schema must live entirely inside the governed
# observation allowlist. Widening the allowlist is a ruling, never an import.
_GOVERNED_COLUMNS = {c.strip().lower() for c in CALIBRATION_OBSERVATION_COLUMNS}
_UNGOVERNED_COLUMNS = sorted(
    (set(MISSION_FIELD_TO_GOVERNED_COLUMN.values()) | set(REQUIRED_DATA_PLANE_COLUMNS))
    - _GOVERNED_COLUMNS
)
if _UNGOVERNED_COLUMNS:  # pragma: no cover - import-time invariant
    raise GovernanceBlock(
        "The data plane schema names columns outside the governed observation allowlist: "
        f"{_UNGOVERNED_COLUMNS}. Extend the allowlist by ruling; a lane schema does not "
        "widen it by import."
    )


#: Admissible as audit/target material only. A subsequent outcome is by
#: construction information from after the observation it accompanies, so it can
#: never be a pregame feature.
NON_FEATURE_COLUMNS: tuple[str, ...] = ("subsequent_outcomes", "actual_margin", "game_result")

SITES: tuple[str, ...] = ("HOME", "AWAY", "NEUTRAL")
RESULTS: tuple[str, ...] = ("WIN", "LOSS", "TIE")


# --- source inventory ---------------------------------------------------------


@dataclass(frozen=True)
class SourceAssessment:
    """One mounted artifact, judged against the historical-observation question."""

    relative_path: str
    content_class: str
    supplies_game_level_observations: bool
    supplies_governed_columns: tuple[str, ...]
    admissible_for_belief_calibration: bool
    reason: str


#: Every historical-source candidate present in the repository, with the reason
#: it does or does not supply game-level observations. FACT in each ``reason``
#: is reproduced from the named artifact; the classification is the assessment.
HISTORICAL_SOURCE_REGISTER: tuple[SourceAssessment, ...] = (
    SourceAssessment(
        relative_path="reference/dynamic_weekly_mc_v3/inputs/2026_FBS_Schedule_LOCKED_v5.xlsx",
        content_class="FORWARD_LOOKING_FIXTURES",
        supplies_game_level_observations=False,
        supplies_governed_columns=("game_id", "week", "event_time", "team", "opponent", "venue"),
        admissible_for_belief_calibration=True,
        reason=(
            "Games sheet carries 743 fixtures for the unplayed 2026 season. Its columns are "
            "game_id, week, date, type, home_team, home_conf, away_team, away_conf, venue, "
            "conference_game, fcs_game, flex_rematch, venue_rule. There is no score column, "
            "so no actual_margin and no game_result can be derived from it."
        ),
    ),
    SourceAssessment(
        relative_path=(
            "reference/dynamic_weekly_mc_v3/inputs/"
            "2026_CFB_2026_Preseason_Unified_Power_Ratings.xlsx"
        ),
        content_class="DERIVED_PRIOR_SEASON_RATINGS",
        supplies_game_level_observations=False,
        supplies_governed_columns=("pregame_team_rating", "pregame_opponent_rating"),
        admissible_for_belief_calibration=True,
        reason=(
            "Master Ratings carries 2025 Baxter Ridge Rating and 2025 Baxter Games for 121 "
            "teams, and Baxter Carryforward Status reads '2025 Baxter fitted rating' for 119 "
            "of them. These are season-level aggregates fitted from 2025 games that are not "
            "mounted anywhere in the repository. A fitted rating cannot be inverted back into "
            "the per-game margins that produced it."
        ),
    ),
    SourceAssessment(
        relative_path="reference/dynamic_weekly_mc_v3/inputs/Model_Parameters_v2_5_APPROVED.xlsx",
        content_class="PARAMETER_AND_GOVERNANCE_REGISTER",
        supplies_game_level_observations=False,
        supplies_governed_columns=(),
        admissible_for_belief_calibration=True,
        reason=(
            "The CALIBRATION sheet records achieved aggregate metrics from the Jul 13-14 "
            "engine run (Margin SD 20.2 against a 16-18 band, OT rate 2.5 against 4-6), which "
            "are simulated engine behaviour rather than observed games. Sheet "
            "17_V3_COMEBACK_RESEARCH labels every comeback target 'Research-only calibration "
            "target; not observed performance' and 'NOT IMPLEMENTED / NOT VALIDATED'."
        ),
    ),
    SourceAssessment(
        relative_path=(
            "reference/dynamic_weekly_mc_v3/inputs/"
            "V2_1_STATIC_CONTROL_2026_CFB_10000_Season_Monte_Carlo.xlsx"
        ),
        content_class="SIMULATED_ENGINE_OUTPUT",
        supplies_game_level_observations=False,
        supplies_governed_columns=(),
        admissible_for_belief_calibration=False,
        reason=(
            "10,000 simulated paths over the 2026 schedule, reported as team-season "
            "probabilities. Simulated margins are engine output; calibrating the engine "
            "against them is circular and is not evidence about any played season."
        ),
    ),
    SourceAssessment(
        relative_path=(
            "reference/dynamic_weekly_mc_v3/inputs/"
            "POWER_CRUNCH_2026_Team_Reconciled_Master_134_v2_2.xlsx"
        ),
        content_class="IDENTITY_RECONCILIATION",
        supplies_game_level_observations=False,
        supplies_governed_columns=("team", "opponent"),
        admissible_for_belief_calibration=True,
        reason=(
            "Reconciled Master resolves 134 schedule entities and their identifiers. It "
            "carries identity and rating-authority columns, and no game rows at all."
        ),
    ),
    SourceAssessment(
        relative_path=(
            "reference/dynamic_weekly_mc_v3/inputs/"
            "2026_TEAM_CANONICAL_MASTER_v2_LLM_GROUNDING.md"
        ),
        content_class="IDENTITY_MASTER",
        supplies_game_level_observations=False,
        supplies_governed_columns=("team", "opponent"),
        admissible_for_belief_calibration=True,
        reason=(
            "Authoritative identity universe of 134 schedule_id values. Supplies the "
            "identity check for team and opponent; carries no game observations."
        ),
    ),
    SourceAssessment(
        relative_path="reference/dynamic_weekly_mc_v3/inputs/2026 Bracket Regime LOCKED.xlsx",
        content_class="POSTSEASON_STRUCTURE",
        supplies_game_level_observations=False,
        supplies_governed_columns=(),
        admissible_for_belief_calibration=True,
        reason="Bracket rules and a projected 2026 field. No played games.",
    ),
    SourceAssessment(
        relative_path="reference/dynamic_weekly_mc_v3/inputs/2026 Playoff Calendar OFFICIAL 2.xlsx",
        content_class="POSTSEASON_CALENDAR",
        supplies_game_level_observations=False,
        supplies_governed_columns=("event_time",),
        admissible_for_belief_calibration=True,
        reason="2026-27 postseason dates and rulings. Future fixtures, no results.",
    ),
    SourceAssessment(
        relative_path=(
            "config/dynamic_weekly_mc_v3/governed/aac_divisions_2026_R2_SUCCESSOR.csv"
        ),
        content_class="CONFERENCE_STRUCTURE",
        supplies_game_level_observations=False,
        supplies_governed_columns=(),
        admissible_for_belief_calibration=True,
        reason="Governed 2026 AAC division assignment. Structure, not observation.",
    ),
    SourceAssessment(
        relative_path="examples/week1_sample_seed.json",
        content_class="MARKET_FLOW_SNAPSHOT",
        supplies_game_level_observations=False,
        supplies_governed_columns=(),
        admissible_for_belief_calibration=False,
        reason=(
            "Five 2026 Week 1 side and total lines from a Covers snapshot. Market lines are "
            "FLOW: they are excluded from BELIEF calibration by doctrine, and they are in any "
            "case forward-looking prices rather than results."
        ),
    ),
)

#: Sources excluded from BELIEF calibration regardless of what they contain.
FLOW_EXCLUDED_SOURCES: tuple[str, ...] = tuple(
    s.relative_path for s in HISTORICAL_SOURCE_REGISTER if not s.admissible_for_belief_calibration
)


def _sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def inventory_historical_sources(root: Path) -> list[dict[str, Any]]:
    """Inventory every historical-source candidate, hashing what is mounted."""
    out: list[dict[str, Any]] = []
    for source in HISTORICAL_SOURCE_REGISTER:
        path = root / source.relative_path
        mounted = path.is_file()
        out.append(
            {
                "relative_path": source.relative_path,
                "mounted": mounted,
                "sha256": _sha256_file(path) if mounted else None,
                "bytes": path.stat().st_size if mounted else None,
                "content_class": source.content_class,
                "supplies_game_level_observations": source.supplies_game_level_observations,
                "supplies_governed_columns": list(source.supplies_governed_columns),
                "admissible_for_belief_calibration": source.admissible_for_belief_calibration,
                "reason": source.reason,
            }
        )
    return out


def assert_flow_source_excluded(relative_path: str) -> None:
    """Refuse a FLOW source presented as BELIEF calibration evidence."""
    if relative_path in FLOW_EXCLUDED_SOURCES:
        raise GovernanceBlock(
            f"Source {relative_path} is not admissible as BELIEF calibration evidence. "
            "Public money and market prices create flow, not belief."
        )


# --- observations -------------------------------------------------------------


@dataclass(frozen=True)
class CalibrationObservation:
    """One normalized, provenance-bound historical observation."""

    game_id: str
    season: int
    week: int
    event_time: datetime
    team: str
    opponent: str
    venue: str
    pregame_team_rating: float
    pregame_opponent_rating: float
    expected_margin: float
    actual_margin: float
    game_result: str
    prior_rating_state: str
    source_provenance: str
    observed_at: datetime
    recorded_at: datetime
    model_version: str
    configuration_version: str
    subsequent_outcomes: str | None = None
    split: str | None = None

    @property
    def is_test_fixture(self) -> bool:
        return self.source_provenance.startswith(TEST_FIXTURE_STAMP)

    @property
    def pairing_key(self) -> tuple[int, int, str, str]:
        """Season, week and the unordered team pair, for mirror detection."""
        a, b = sorted((self.team, self.opponent))
        return (self.season, self.week, a, b)

    def as_dict(self) -> dict[str, Any]:
        return {
            "game_id": self.game_id,
            "season": self.season,
            "week": self.week,
            "event_time": self.event_time.isoformat(),
            "team": self.team,
            "opponent": self.opponent,
            "venue": self.venue,
            "pregame_team_rating": self.pregame_team_rating,
            "pregame_opponent_rating": self.pregame_opponent_rating,
            "expected_margin": self.expected_margin,
            "actual_margin": self.actual_margin,
            "game_result": self.game_result,
            "prior_rating_state": self.prior_rating_state,
            "subsequent_outcomes": self.subsequent_outcomes,
            "source_provenance": self.source_provenance,
            "observed_at": self.observed_at.isoformat(),
            "recorded_at": self.recorded_at.isoformat(),
            "model_version": self.model_version,
            "configuration_version": self.configuration_version,
            "split": self.split,
        }


def _require_utc(value: str, field_name: str, game_id: str) -> datetime:
    raw = value.strip()
    if not raw:
        raise InputValidationError(f"Observation {game_id} has empty {field_name}")
    text = raw[:-1] + "+00:00" if raw.endswith("Z") else raw
    try:
        parsed = datetime.fromisoformat(text)
    except ValueError:
        raise InputValidationError(
            f"Observation {game_id} has unparseable {field_name} {value!r}; ISO-8601 required"
        ) from None
    if parsed.tzinfo is None:
        raise InputValidationError(
            f"Observation {game_id} has naive {field_name} {value!r}. A timestamp without an "
            "offset cannot be ordered against another source's timestamps."
        )
    return parsed.astimezone(timezone.utc)


def _require_number(value: str, field_name: str, game_id: str) -> float:
    try:
        return float(str(value).strip())
    except (TypeError, ValueError):
        raise InputValidationError(
            f"Observation {game_id} has non-numeric {field_name} {value!r}"
        ) from None


def _require_int(value: str, field_name: str, game_id: str) -> int:
    text = str(value).strip()
    try:
        return int(text)
    except (TypeError, ValueError):
        raise InputValidationError(
            f"Observation {game_id} has non-integer {field_name} {value!r}"
        ) from None


def normalize_row(row: dict[str, Any]) -> CalibrationObservation:
    """Normalize one raw row into a typed observation. Fail-closed on every field."""
    clean = {str(k).strip().lower(): ("" if v is None else str(v).strip()) for k, v in row.items()}
    game_id = clean.get("game_id", "")
    if not game_id:
        raise InputValidationError("Observation row carries no game_id")

    missing = [c for c in REQUIRED_DATA_PLANE_COLUMNS if not clean.get(c)]
    if missing:
        raise InputValidationError(
            f"Observation {game_id} is missing required values for {missing}. "
            f"{MISSING_CALIBRATION_EVIDENCE}: a partial observation is not completed by default."
        )

    venue = clean["venue"].upper()
    if venue not in SITES:
        raise InputValidationError(
            f"Observation {game_id} has site {clean['venue']!r}; expected one of {list(SITES)}"
        )
    result = clean["game_result"].upper()
    if result not in RESULTS:
        raise InputValidationError(
            f"Observation {game_id} has result {clean['game_result']!r}; "
            f"expected one of {list(RESULTS)}"
        )

    subsequent = clean.get("subsequent_outcomes") or None
    split = (clean.get("split") or "").lower() or None
    if split is not None and split not in DATA_SPLITS:
        raise InputValidationError(
            f"Observation {game_id} has split {split!r}; expected one of {list(DATA_SPLITS)}"
        )

    return CalibrationObservation(
        game_id=game_id,
        season=_require_int(clean["season"], "season", game_id),
        week=_require_int(clean["week"], "week", game_id),
        event_time=_require_utc(clean["event_time"], "event_time", game_id),
        team=clean["team"],
        opponent=clean["opponent"],
        venue=venue,
        pregame_team_rating=_require_number(clean["pregame_team_rating"], "pregame_team_rating", game_id),
        pregame_opponent_rating=_require_number(
            clean["pregame_opponent_rating"], "pregame_opponent_rating", game_id
        ),
        expected_margin=_require_number(clean["expected_margin"], "expected_margin", game_id),
        actual_margin=_require_number(clean["actual_margin"], "actual_margin", game_id),
        game_result=result,
        prior_rating_state=clean["prior_rating_state"],
        source_provenance=clean["source_provenance"],
        observed_at=_require_utc(clean["observed_at"], "observed_at", game_id),
        recorded_at=_require_utc(clean["recorded_at"], "recorded_at", game_id),
        model_version=clean["model_version"],
        configuration_version=clean["configuration_version"],
        subsequent_outcomes=subsequent,
        split=split,
    )


def _raw_rows(path: Path, fmt: str) -> list[dict[str, Any]]:
    text = path.read_text(encoding="utf-8")
    if fmt in ("csv", "tsv"):
        delimiter = "," if fmt == "csv" else "\t"
        reader = csv.DictReader(io.StringIO(text, newline=""), delimiter=delimiter)
        return [dict(r) for r in reader]
    payload = json.loads(text)
    rows = payload.get("observations") if isinstance(payload, dict) else payload
    return [dict(r) for r in rows]


def load_observations(
    path: Path, dataset_id: str, *, fmt: str | None = None
) -> tuple[list[CalibrationObservation], dict[str, Any]]:
    """Load and normalize an observation set, returning it with its provenance.

    Admissibility is delegated to :func:`.calibration.register_dataset` first, so
    a set carrying a forbidden signal or an ungoverned column is refused before a
    single row is parsed.
    """
    registered = register_dataset(path, dataset_id, fmt=fmt)
    declared = (fmt or path.suffix.lstrip(".")).strip().lower()

    columns = {c.strip().lower() for c in registered.columns}
    missing = sorted(c for c in REQUIRED_DATA_PLANE_COLUMNS if c not in columns)
    if missing:
        raise GovernanceBlock(
            f"Calibration dataset {dataset_id} is missing data-plane columns: {missing}. "
            f"{MISSING_CALIBRATION_EVIDENCE}: the lane schema is not satisfied by this set."
        )

    observations = [normalize_row(r) for r in _raw_rows(path, declared)]
    if not observations:
        raise GovernanceBlock(
            f"Calibration dataset {dataset_id} parsed to zero observations. "
            f"{BLOCKED_ON_CALIBRATION_DATA}."
        )
    provenance = {
        "dataset_id": registered.dataset_id,
        "path": str(path),
        "sha256": registered.sha256,
        "declared_format": declared,
        "columns": list(registered.columns),
        "observation_count": len(observations),
        "contains_test_fixture_rows": any(o.is_test_fixture for o in observations),
    }
    return observations, provenance


def assert_not_test_fixture(
    observations: Sequence[CalibrationObservation], dataset_id: str
) -> None:
    """Refuse a synthetic set presented as governed calibration evidence."""
    stamped = sorted({o.game_id for o in observations if o.is_test_fixture})
    if stamped:
        raise GovernanceBlock(
            f"Calibration dataset {dataset_id} carries TEST_FIXTURE observations {stamped}. "
            "A synthetic fixture exercises the data plane and never becomes evidence."
        )


# --- deterministic checks -----------------------------------------------------

BLOCK = "BLOCK"
WARN = "WARN"


@dataclass(frozen=True)
class Finding:
    """One deterministic check result."""

    check: str
    code: str
    severity: str
    game_id: str | None
    detail: str

    def as_dict(self) -> dict[str, Any]:
        return {
            "check": self.check,
            "code": self.code,
            "severity": self.severity,
            "game_id": self.game_id,
            "detail": self.detail,
        }


def _sorted_findings(findings: Iterable[Finding]) -> list[Finding]:
    return sorted(findings, key=lambda f: (f.check, f.code, f.game_id or "", f.detail))


def check_identity(
    observations: Sequence[CalibrationObservation], known_team_ids: set[str]
) -> list[Finding]:
    """Every team and opponent must resolve in the canonical universe, exactly.

    An unresolved identifier blocks. It is never matched by similarity, because a
    near-match is how one team's season quietly becomes another team's evidence.
    """
    findings: list[Finding] = []
    for o in observations:
        for role, value in (("team", o.team), ("opponent", o.opponent)):
            if value not in known_team_ids:
                findings.append(
                    Finding(
                        "identity",
                        "UNRESOLVED_TEAM_IDENTIFIER",
                        BLOCK,
                        o.game_id,
                        f"{role}={value!r} is not a canonical schedule_id; no guess is made",
                    )
                )
        if o.team == o.opponent:
            findings.append(
                Finding(
                    "identity",
                    "SELF_OPPONENT",
                    BLOCK,
                    o.game_id,
                    f"team and opponent are both {o.team!r}",
                )
            )
        declared = o.game_result
        derived = "TIE" if o.actual_margin == 0 else ("WIN" if o.actual_margin > 0 else "LOSS")
        if declared != derived:
            findings.append(
                Finding(
                    "identity",
                    "RESULT_MARGIN_DISAGREEMENT",
                    BLOCK,
                    o.game_id,
                    f"result {declared} disagrees with actual_margin {o.actual_margin}",
                )
            )
    return _sorted_findings(findings)


def check_chronology(observations: Sequence[CalibrationObservation]) -> list[Finding]:
    """Timestamps must order as event -> observed -> recorded, and weeks must rise."""
    findings: list[Finding] = []
    for o in observations:
        if o.observed_at < o.event_time:
            findings.append(
                Finding(
                    "chronology",
                    "OBSERVED_BEFORE_EVENT",
                    BLOCK,
                    o.game_id,
                    f"observed_at {o.observed_at.isoformat()} precedes event_time "
                    f"{o.event_time.isoformat()}: the outcome cannot be known before the game",
                )
            )
        if o.recorded_at < o.observed_at:
            findings.append(
                Finding(
                    "chronology",
                    "RECORDED_BEFORE_OBSERVED",
                    BLOCK,
                    o.game_id,
                    f"recorded_at {o.recorded_at.isoformat()} precedes observed_at "
                    f"{o.observed_at.isoformat()}",
                )
            )
        if o.week < 1:
            findings.append(
                Finding("chronology", "NON_POSITIVE_WEEK", BLOCK, o.game_id, f"week={o.week}")
            )

    by_season: dict[int, list[CalibrationObservation]] = {}
    for o in observations:
        by_season.setdefault(o.season, []).append(o)
    for season, rows in sorted(by_season.items()):
        ordered = sorted(rows, key=lambda r: (r.event_time, r.game_id))
        weeks_seen: dict[int, datetime] = {}
        for o in ordered:
            weeks_seen.setdefault(o.week, o.event_time)
        previous_week = None
        previous_start = None
        for week, start in sorted(weeks_seen.items()):
            if previous_start is not None and start < previous_start:
                findings.append(
                    Finding(
                        "chronology",
                        "WEEK_EVENT_TIME_INVERSION",
                        BLOCK,
                        None,
                        f"season {season} week {week} starts {start.isoformat()}, before "
                        f"week {previous_week} at {previous_start.isoformat()}",
                    )
                )
            previous_week, previous_start = week, start
    return _sorted_findings(findings)


def check_duplicates(observations: Sequence[CalibrationObservation]) -> list[Finding]:
    """Detect repeated game_ids and both-perspective mirrors of the same fixture."""
    findings: list[Finding] = []
    by_id: dict[str, list[CalibrationObservation]] = {}
    for o in observations:
        by_id.setdefault(o.game_id, []).append(o)
    for game_id, rows in sorted(by_id.items()):
        if len(rows) > 1:
            findings.append(
                Finding(
                    "duplicates",
                    "DUPLICATE_GAME_ID",
                    BLOCK,
                    game_id,
                    f"{len(rows)} rows share game_id {game_id!r}",
                )
            )

    by_pair: dict[tuple[int, int, str, str], list[CalibrationObservation]] = {}
    for o in observations:
        by_pair.setdefault(o.pairing_key, []).append(o)
    for key, rows in sorted(by_pair.items()):
        if len(rows) > 1:
            season, week, a, b = key
            findings.append(
                Finding(
                    "duplicates",
                    "MIRRORED_OBSERVATION",
                    BLOCK,
                    sorted(r.game_id for r in rows)[0],
                    f"season {season} week {week} pair {a}/{b} appears {len(rows)} times; "
                    "both perspectives of one fixture double-count it",
                )
            )
    return _sorted_findings(findings)


def check_leakage(observations: Sequence[CalibrationObservation]) -> list[Finding]:
    """Refuse observations whose pregame state could have seen the outcome.

    ``prior_rating_state`` is checked when it carries an ``@<iso8601>`` as-of
    suffix. When it does not, the state is unverifiable rather than clean, and
    that is reported instead of assumed away.
    """
    findings: list[Finding] = []
    for o in observations:
        state = o.prior_rating_state
        if "@" in state:
            _, _, stamp = state.rpartition("@")
            try:
                as_of = _require_utc(stamp, "prior_rating_state", o.game_id)
            except InputValidationError:
                findings.append(
                    Finding(
                        "leakage",
                        "UNPARSEABLE_PRIOR_RATING_STATE_AS_OF",
                        BLOCK,
                        o.game_id,
                        f"prior_rating_state {state!r} carries an unparseable as-of stamp",
                    )
                )
            else:
                if as_of > o.event_time:
                    findings.append(
                        Finding(
                            "leakage",
                            "PRIOR_RATING_STATE_AFTER_EVENT",
                            BLOCK,
                            o.game_id,
                            f"prior_rating_state as-of {as_of.isoformat()} is later than "
                            f"event_time {o.event_time.isoformat()}",
                        )
                    )
        else:
            findings.append(
                Finding(
                    "leakage",
                    "UNVERIFIABLE_PRIOR_RATING_STATE",
                    WARN,
                    o.game_id,
                    f"prior_rating_state {state!r} carries no @<iso8601> as-of stamp, so its "
                    "freedom from hindsight cannot be established from the row alone",
                )
            )
        if o.subsequent_outcomes and o.split == "training":
            findings.append(
                Finding(
                    "leakage",
                    "SUBSEQUENT_OUTCOMES_IN_TRAINING_ROW",
                    WARN,
                    o.game_id,
                    "subsequent_outcomes is admissible as audit/target material only and must "
                    "not be consumed as a pregame feature",
                )
            )
    return _sorted_findings(findings)


def check_split_separation(observations: Sequence[CalibrationObservation]) -> list[Finding]:
    """Each game lands in exactly one split, and the splits do not cross in time."""
    findings: list[Finding] = []
    assigned = [o for o in observations if o.split is not None]
    if len(assigned) != len(observations):
        findings.append(
            Finding(
                "split",
                "UNASSIGNED_OBSERVATIONS",
                BLOCK,
                None,
                f"{len(observations) - len(assigned)} observations carry no split",
            )
        )
        return _sorted_findings(findings)

    by_game: dict[str, set[str]] = {}
    for o in assigned:
        by_game.setdefault(o.game_id, set()).add(str(o.split))
    for game_id, splits in sorted(by_game.items()):
        if len(splits) > 1:
            findings.append(
                Finding(
                    "split",
                    "OBSERVATION_IN_MULTIPLE_SPLITS",
                    BLOCK,
                    game_id,
                    f"game {game_id} appears in {sorted(splits)}",
                )
            )

    bounds: dict[str, tuple[datetime, datetime]] = {}
    for name in DATA_SPLITS:
        rows = [o for o in assigned if o.split == name]
        if rows:
            times = [o.event_time for o in rows]
            bounds[name] = (min(times), max(times))
    order = [n for n in DATA_SPLITS if n in bounds]
    for earlier, later in zip(order, order[1:]):
        if bounds[earlier][1] > bounds[later][0]:
            findings.append(
                Finding(
                    "split",
                    "SPLIT_CHRONOLOGY_OVERLAP",
                    BLOCK,
                    None,
                    f"{earlier} extends to {bounds[earlier][1].isoformat()}, past the start of "
                    f"{later} at {bounds[later][0].isoformat()}; a later game must not train a "
                    "model evaluated on an earlier one",
                )
            )
    return _sorted_findings(findings)


def run_all_checks(
    observations: Sequence[CalibrationObservation], known_team_ids: set[str]
) -> dict[str, Any]:
    """Run every deterministic check and report the aggregate disposition."""
    findings = (
        check_identity(observations, known_team_ids)
        + check_chronology(observations)
        + check_duplicates(observations)
        + check_leakage(observations)
        + check_split_separation(observations)
    )
    findings = _sorted_findings(findings)
    blocking = [f for f in findings if f.severity == BLOCK]
    return {
        "observation_count": len(observations),
        "findings": [f.as_dict() for f in findings],
        "blocking_findings": [f.as_dict() for f in blocking],
        "admissible": not blocking,
        "disposition": "DATA_PLANE_CHECKS_PASSED" if not blocking else "DATA_PLANE_CHECKS_BLOCKED",
    }


# --- deterministic splitting --------------------------------------------------


@dataclass(frozen=True)
class SplitPolicy:
    """How observations are divided. Deterministic, and explicitly not governed.

    Chronological rather than hashed: a weekly rerating model that trains on week
    14 to be evaluated on week 3 has seen the future, and a hashed split does
    exactly that while looking neutral. No ruling names a split method, so this
    policy carries its own ungoverned status rather than borrowing authority from
    ruling R2-CAL-OBJECTIVE, which names only the three split *names*.
    """

    policy_id: str = "CHRONOLOGICAL_60_20_20"
    training_fraction: float = 0.60
    validation_fraction: float = 0.20
    status: str = "NOT_GOVERNED_PENDING_RULING"

    def __post_init__(self) -> None:
        if not 0.0 < self.training_fraction < 1.0:
            raise InputValidationError(f"training_fraction out of range: {self.training_fraction}")
        if not 0.0 < self.validation_fraction < 1.0:
            raise InputValidationError(
                f"validation_fraction out of range: {self.validation_fraction}"
            )
        if self.training_fraction + self.validation_fraction >= 1.0:
            raise InputValidationError(
                "training and validation fractions leave no holdout: "
                f"{self.training_fraction} + {self.validation_fraction}"
            )

    def as_dict(self) -> dict[str, Any]:
        return {
            "policy_id": self.policy_id,
            "training_fraction": self.training_fraction,
            "validation_fraction": self.validation_fraction,
            "holdout_fraction": round(1.0 - self.training_fraction - self.validation_fraction, 10),
            "method": "CHRONOLOGICAL",
            "tie_break": "game_id ascending",
            "randomised": False,
            "status": self.status,
            "governed": False,
        }


DEFAULT_SPLIT_POLICY = SplitPolicy()


def assign_deterministic_splits(
    observations: Sequence[CalibrationObservation],
    policy: SplitPolicy = DEFAULT_SPLIT_POLICY,
) -> list[CalibrationObservation]:
    """Assign training/validation/holdout chronologically and reproducibly.

    Rows are ordered by ``(event_time, game_id)``, so the assignment depends on
    the data alone: no seed, no clock, no iteration order. Every split must
    receive at least one observation, because an empty holdout would let a
    promotion claim holdout evidence that does not exist.
    """
    if not observations:
        raise GovernanceBlock(
            f"Cannot split an empty observation set. {BLOCKED_ON_CALIBRATION_DATA}."
        )
    total = len(observations)
    if total < len(DATA_SPLITS):
        raise GovernanceBlock(
            f"{total} observations cannot fill {len(DATA_SPLITS)} splits. "
            f"{BLOCKED_ON_CALIBRATION_DATA}: an empty split is not a split."
        )
    ordered = sorted(observations, key=lambda o: (o.event_time, o.game_id))
    train_end = int(total * policy.training_fraction)
    validation_end = train_end + int(total * policy.validation_fraction)
    train_end = max(1, min(train_end, total - 2))
    validation_end = max(train_end + 1, min(validation_end, total - 1))

    out: list[CalibrationObservation] = []
    for index, observation in enumerate(ordered):
        if index < train_end:
            split = "training"
        elif index < validation_end:
            split = "validation"
        else:
            split = "holdout"
        out.append(replace(observation, split=split))
    return out


# --- missing-evidence manifest ------------------------------------------------
#
# The mounted artifacts cannot supply a historical observation set, so the
# deliverable of this lane is a precise statement of what is absent rather than
# a dataset assembled to fill the gap.

#: Where the loader reads an observation set from. DERIVED: this is the path
#: this module reads, not a governed artifact name. Mounting a file here is a
#: data step; it does not by itself make the file authoritative.
EXPECTED_OBSERVATION_SET = "reference/dynamic_weekly_mc_v3/inputs/calibration/historical_game_observations.csv"

#: Artifacts governance already records as named but not mounted. FACT, quoted
#: from V3_GOVERNANCE_STATUS_R4.json; not re-derived and not extended here.
GOVERNANCE_RECORDED_ABSENT_ARTIFACTS: tuple[dict[str, Any], ...] = (
    {
        "artifact": "2026_Board_I-K_CANONICAL_APPROVED_R1_REISSUE(3).xlsx",
        "required_sha256": "6b4cec1e48b34cb9eca5c224f5ef8750cca5acc40ecfd0bc42ed62976cbd4c9a",
        "blocker": "inputs.board_of_record_i_k",
        "recorded_in": "V3_GOVERNANCE_STATUS_R4.json!board_of_record",
        "relevance_to_calibration": (
            "Board of record for rating authority. Not an observation source; listed so a "
            "reader does not mistake its absence for the calibration gap or the reverse."
        ),
    },
    {
        "artifact": "srs_spec.md, compute_srs.py, srs_2025.csv, srs_validation_2025.txt",
        "required_sha256": None,
        "blocker": None,
        "recorded_in": "V3_GOVERNANCE_STATUS_R4.json!srs.evidence_searched_for",
        "relevance_to_calibration": (
            "SRS is a named independent witness under ruling R2-CAL-OBJECTIVE. srs_2025.csv "
            "would be a 2025-season artifact, but a validation file of season ratings is "
            "still not the game-level observation set this lane requires."
        ),
    },
)

#: Fields the mounted evidence cannot supply at game level, with the reason.
UNSUPPLIED_REQUIRED_FIELDS: tuple[str, ...] = (
    "actual_margin",
    "game_result",
    "expected_margin",
    "pregame_team_rating",
    "pregame_opponent_rating",
    "prior_rating_state",
    "observed_at",
    "recorded_at",
    "source_provenance",
    "model_version",
    "configuration_version",
)


def summarize_absent_prior_season_games(ratings_xlsx: Path) -> dict[str, Any]:
    """Quantify the 2025 games the mounted ratings were fitted on but do not carry.

    ``2025 Baxter Games`` records how many games each team's fitted rating drew
    on. Summing it gives the number of team-game *sides* whose per-game record is
    absent; the game count is at least half that, and only exactly half if every
    opponent were inside the 121-team set, which the FCS schedule rows rule out.
    """
    from openpyxl import load_workbook

    workbook = load_workbook(ratings_xlsx, read_only=True, data_only=True)
    try:
        sheet = workbook["Master Ratings"]
        header = [cell.value for cell in next(sheet.iter_rows(min_row=4, max_row=4))]
        index = {str(name): position for position, name in enumerate(header)}
        rows = [
            row
            for row in sheet.iter_rows(min_row=5, values_only=True)
            if row[index["Schedule ID"]] is not None
        ]
        games_column = index["2025 Baxter Games"]
        counts = [int(row[games_column] or 0) for row in rows]
    finally:
        workbook.close()

    sides = sum(counts)
    return {
        "teams_with_fitted_2025_rating": sum(1 for c in counts if c > 0),
        "teams_without_fitted_2025_rating": sum(1 for c in counts if c == 0),
        "team_game_sides_recorded": sides,
        "implied_game_count_lower_bound": -(-sides // 2),
        "per_game_rows_mounted": 0,
        "note": (
            "2025 Baxter Games is a count, not a record. No mounted artifact carries the "
            "individual 2025 games, their margins, their dates or their pregame ratings."
        ),
    }


def build_missing_calibration_evidence_manifest(root: Path) -> dict[str, Any]:
    """Build the machine-readable manifest of what calibration evidence is absent.

    The manifest is a pure function of repository contents. It carries no
    generation timestamp on purpose: a clock would make two runs over identical
    evidence produce different artifacts, and the whole point is that the same
    evidence always yields the same answer.
    """
    from .config import V3Config

    config = V3Config.from_json(root / "config/dynamic_weekly_mc_v3/v3_experimental.json")
    sources = inventory_historical_sources(root)
    supplying = [s for s in sources if s["supplies_game_level_observations"]]
    expected_path = root / EXPECTED_OBSERVATION_SET

    prior_season = summarize_absent_prior_season_games(
        root / "reference/dynamic_weekly_mc_v3/inputs/"
        "2026_CFB_2026_Preseason_Unified_Power_Ratings.xlsx"
    )

    field_status = {}
    for mission_field, column in sorted(MISSION_FIELD_TO_GOVERNED_COLUMN.items()):
        partial = sorted(
            s["relative_path"] for s in sources if column in s["supplies_governed_columns"]
        )
        field_status[mission_field] = {
            "governed_column": column,
            "required": column in REQUIRED_DATA_PLANE_COLUMNS,
            "available_at_game_level": False,
            "partially_supplied_by": partial,
            "status": "PARTIAL_NON_OBSERVATIONAL" if partial else "ABSENT",
        }

    return {
        "code": MISSING_CALIBRATION_EVIDENCE,
        "artifact": "MISSING_CALIBRATION_EVIDENCE.json",
        "scope": "HISTORICAL_CALIBRATION_DATA_PLANE",
        "model_version": config.model_version,
        "configuration_version": config.configuration_version,
        "disposition": BLOCKED_ON_CALIBRATION_DATA,
        "summary": (
            "No mounted artifact carries game-level historical observations. Every source is "
            "either a 2026 fixture list for an unplayed season, a season-level rating derived "
            "from games that are not mounted, simulated engine output, identity or structure "
            "reference, or market flow excluded from belief calibration."
        ),
        "observation_sets_mounted": 0,
        "sources_supplying_game_level_observations": len(supplying),
        "missing_files": [
            {
                "path": EXPECTED_OBSERVATION_SET,
                "exists": expected_path.exists(),
                "provenance_class": "DERIVED",
                "role": "historical game-level observation set the data plane loads",
                "required_columns": list(REQUIRED_DATA_PLANE_COLUMNS),
                "optional_columns": ["subsequent_outcomes", "split"],
                "accepted_formats": ["csv", "tsv", "json"],
                "note": (
                    "Path naming is DERIVED from this module's loader, not from a governed "
                    "artifact register. Mounting a file here does not confer authority."
                ),
            }
        ],
        "missing_fields": {
            "required_but_unsupplied_at_game_level": list(UNSUPPLIED_REQUIRED_FIELDS),
            "per_field": field_status,
        },
        "governance_recorded_absent_artifacts": [
            dict(entry) for entry in GOVERNANCE_RECORDED_ABSENT_ARTIFACTS
        ],
        "prior_season_fit_without_underlying_games": prior_season,
        "source_inventory": sources,
        "flow_excluded_sources": list(FLOW_EXCLUDED_SOURCES),
        "flow_exclusion_basis": (
            "Public money and market prices create flow, not belief. Excluded from BELIEF "
            "calibration at source registration and again at column admissibility."
        ),
        "split_policy": DEFAULT_SPLIT_POLICY.as_dict(),
        "unresolved_calibration_fields": list(CALIBRATION_FIELDS),
        "calibration_blockers_retired_by_this_lane": [],
        "canonical_values_written": [],
        "data_plane_ready": True,
        "data_plane_ready_note": (
            "Loaders, normalization, identity, chronology, duplicate, leakage and split checks "
            "are implemented and tested against TEST_FIXTURE data. Working plumbing is not "
            "evidence: no calibration blocker is retired and no coefficient is proposed."
        ),
        "required_to_clear": [
            "Mount a governed game-level historical observation set carrying every required "
            "column, with source artifact provenance and hashes per row.",
            "Obtain a ruling fixing the train/validation/holdout split policy; the "
            "chronological default here is deterministic but not governed.",
            "Run the experimental harness against the primary criterion named by ruling "
            "R2-CAL-OBJECTIVE, reporting Colley and SRS independently.",
            "Promote only by explicit human approval token under a named authority.",
        ],
    }


def emit_missing_calibration_evidence_manifest(root: Path, destination: Path) -> dict[str, Any]:
    """Write the manifest deterministically and return it."""
    manifest = build_missing_calibration_evidence_manifest(root)
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(
        json.dumps(manifest, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )
    return manifest


def canonical_team_universe(canonical_md: Path) -> set[str]:
    """The 134 canonical schedule_id values, for identity checks."""
    from .inputs import load_canonical_team_index

    return set(load_canonical_team_index(canonical_md))
