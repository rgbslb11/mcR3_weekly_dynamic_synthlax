from __future__ import annotations

import csv
import hashlib
import io
import re
from pathlib import Path

from openpyxl import load_workbook

from .errors import InputValidationError
from .models import ScheduledGame, Team
from .provenance import (
    BINARY_MISMATCH,
    CERTIFIED_GAMES_CONTENT_SHA256,
    CONTENT_MISMATCH,
    GOVERNED_SCHEDULE_BINARY_SHA256,
    ScheduleProvenance,
)


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _bool_yn(value: object) -> bool:
    return str(value or "").strip().upper() == "Y"


def load_canonical_team_index(markdown_path: Path) -> dict[str, dict[str, str | None]]:
    text = markdown_path.read_text(encoding="utf-8")
    start = text.find("## Canonical Team Index — All 134 Entities")
    if start < 0:
        raise InputValidationError("Canonical team index section not found")
    section = text[start:]
    rows: dict[str, dict[str, str | None]] = {}
    pattern = re.compile(
        r"^\|\s*\d+\s*\|\s*`([^`]+)`\s*\|\s*([^|]+?)\s*\|\s*`([^`]+)`\s*\|\s*`([^`]+)`\s*\|\s*`([^`]+)`\s*\|\s*(.*?)\s*\|\s*(.*?)\s*\|\s*(.*?)\s*\|$",
        re.MULTILINE,
    )
    for m in pattern.finditer(section):
        schedule_id, team_name, abbreviated, scope, division, conference, master_id, name_id = m.groups()
        def clean(v: str) -> str | None:
            v = v.strip().replace("`", "")
            return None if v == "<blank>" else v
        rows[schedule_id] = {
            "schedule_id": schedule_id,
            "team_name": team_name.strip(),
            "abbreviated_name": abbreviated,
            "entity_scope": scope,
            "division": division,
            "conference": clean(conference),
            "master_team_id": clean(master_id),
            "name_id": clean(name_id),
        }
    if len(rows) != 134:
        raise InputValidationError(f"Expected 134 canonical entities, found {len(rows)}")
    return rows


def load_preseason_ratings(path: Path) -> dict[str, dict[str, float | str]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Master Ratings"]
    header = [c.value for c in next(ws.iter_rows(min_row=4, max_row=4))]
    idx = {str(name): i for i, name in enumerate(header)}
    required = ["Schedule ID", "Unified Master Power Index", "Unified Neutral-Field Points"]
    missing = [x for x in required if x not in idx]
    if missing:
        raise InputValidationError(f"Unified ratings missing columns: {missing}")
    out: dict[str, dict[str, float | str]] = {}
    for row in ws.iter_rows(min_row=5, values_only=True):
        sid = row[idx["Schedule ID"]]
        if sid is None:
            continue
        out[str(sid)] = {
            "team_name": str(row[idx["Team"]]),
            "conference": str(row[idx["Conference"]]),
            "power_index": float(row[idx["Unified Master Power Index"]]),
            "strength_points": float(row[idx["Unified Neutral-Field Points"]]),
        }
    if len(out) != 121:
        raise InputValidationError(f"Expected 121 FBS preseason ratings, found {len(out)}")
    return out


def load_fcs_operator_rows(path: Path) -> dict[str, dict[str, float | str]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Reconciled Master"]
    rows = ws.iter_rows(values_only=True)
    header = [str(x) if x is not None else "" for x in next(rows)]
    idx = {name: i for i, name in enumerate(header)}
    out: dict[str, dict[str, float | str]] = {}
    for row in rows:
        if row[idx["entity_scope"]] != "SCHEDULE_ONLY_FCS":
            continue
        sid = str(row[idx["schedule_id"]])
        out[sid] = {
            "sim_rating": float(row[idx["sim_rating"]]),
            "board_power_H_equivalent": float(row[idx["board_power_H_equivalent"]]),
            "rating_source": str(row[idx["rating_source"]]),
            "rating_ruling": str(row[idx["rating_ruling"]]),
        }
    if len(out) != 13:
        raise InputValidationError(f"Expected 13 FCS operator rows, found {len(out)}")
    return out


def load_teams(canonical_md: Path, ratings_xlsx: Path) -> dict[str, Team]:
    canonical = load_canonical_team_index(canonical_md)
    ratings = load_preseason_ratings(ratings_xlsx)
    teams: dict[str, Team] = {}
    for sid, row in canonical.items():
        if row["entity_scope"] == "FBS_MEMBER":
            if sid not in ratings:
                raise InputValidationError(f"FBS team {sid} missing preseason rating")
            r = ratings[sid]
            teams[sid] = Team(
                schedule_id=sid,
                team_name=str(row["team_name"]),
                conference=row["conference"],
                entity_scope="FBS_MEMBER",
                championship_eligible=True,
                preseason_power_index=float(r["power_index"]),
                preseason_strength_points=float(r["strength_points"]),
                hfa_modifier=1.0,
            )
        else:
            teams[sid] = Team(
                schedule_id=sid,
                team_name=str(row["team_name"]),
                conference=row["conference"],
                entity_scope="SCHEDULE_ONLY_FCS",
                championship_eligible=False,
                preseason_power_index=None,
                preseason_strength_points=None,
                hfa_modifier=None,
            )
    return teams


def load_schedule(path: Path) -> list[ScheduledGame]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Games"]
    rows = ws.iter_rows(values_only=True)
    header = [str(x) if x is not None else "" for x in next(rows)]
    idx = {name: i for i, name in enumerate(header)}
    games: list[ScheduledGame] = []
    for row in rows:
        def get(name: str):
            i = idx[name]
            return row[i] if i < len(row) else None
        if not row or get("game_id") is None:
            continue
        week_s = str(get("week"))
        games.append(
            ScheduledGame(
                game_id=str(get("game_id")),
                week=int(week_s.lstrip("W")),
                date=str(get("date")),
                game_type=str(get("type")),
                home_team=None if get("home_team") is None else str(get("home_team")),
                away_team=None if get("away_team") is None else str(get("away_team")),
                home_conf=None if get("home_conf") is None else str(get("home_conf")),
                away_conf=None if get("away_conf") is None else str(get("away_conf")),
                venue=str(get("venue")),
                conference_game=_bool_yn(get("conference_game")),
                fcs_game=_bool_yn(get("fcs_game")),
                flex_rematch=_bool_yn(get("flex_rematch")),
                venue_rule=None if get("venue_rule") is None else str(get("venue_rule")),
            )
        )
    return games


def schedule_content_hash(path: Path) -> str:
    """Reproduce the certified Games-sheet content hash.

    The canonical method is documented on the schedule workbook's Certification
    sheet: "SHA-256 of Games sheet serialized as UTF-8 CSV with header, LF line
    endings, and blank for null."

    A row that ends short of the header width has trailing *null* cells, so the
    canonical serialization emits them as blanks. 735 of the 743 data rows omit
    the trailing ``venue_rule`` cell; writing them ragged produces a different
    CSV and therefore a different digest. Every row is padded to the header
    width before serialization so the reproduction matches the certified value.
    """
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Games"]
    rows = [list(r) for r in ws.iter_rows(values_only=True) if any(v is not None for v in r)]
    if not rows:
        raise InputValidationError("Games sheet is empty")
    width = len(rows[0])
    sio = io.StringIO(newline="")
    writer = csv.writer(sio, lineterminator="\n")
    for row in rows:
        padded = (row + [None] * (width - len(row))) if len(row) < width else row[:width]
        writer.writerow(["" if v is None else v for v in padded])
    return hashlib.sha256(sio.getvalue().encode("utf-8")).hexdigest()


def validate_schedule(path: Path, teams: dict[str, Team]) -> dict[str, object]:
    games = load_schedule(path)
    regular = [g for g in games if g.game_type == "REG"]
    ccg = [g for g in games if g.game_type == "CCG"]
    unknown = sorted({t for g in regular for t in (g.home_team, g.away_team) if t and t not in teams})
    if unknown:
        raise InputValidationError(f"Schedule references unknown canonical IDs: {unknown}")
    if len(games) != 743 or len(regular) != 736 or len(ccg) != 7:
        raise InputValidationError(
            f"Schedule count mismatch total={len(games)} regular={len(regular)} ccg={len(ccg)}"
        )
    content_hash = schedule_content_hash(path)
    certified_content_hash = CERTIFIED_GAMES_CONTENT_SHA256
    governed_binary_hash = GOVERNED_SCHEDULE_BINARY_SHA256
    actual_binary_hash = sha256_file(path)
    provenance = ScheduleProvenance(
        content_sha256_reproduced=content_hash,
        content_sha256_certified=certified_content_hash,
        binary_sha256_actual=actual_binary_hash,
        binary_sha256_governed=governed_binary_hash,
    )
    # The full observation is preserved; only the *blocking* subset narrows under
    # ruling R2-SCHED-V5-AUTH.
    anomalies = provenance.anomalies()
    blocking = provenance.blocking_anomalies()
    return {
        "games_total": len(games),
        "regular_games": len(regular),
        "ccg_templates": len(ccg),
        "schedule_games_sha256_reproduced": content_hash,
        "schedule_games_sha256_certified": certified_content_hash,
        "schedule_binary_sha256_actual": actual_binary_hash,
        "schedule_binary_sha256_governed": governed_binary_hash,
        "provenance_anomalies": anomalies,
        "blocking_provenance_anomalies": blocking,
        "governed_binary_disposition": provenance.governed_binary_disposition(),
        "schedule_v5_authority_ruling": "R2-SCHED-V5-AUTH",
        "weeks": sorted({g.week for g in regular}),
    }
