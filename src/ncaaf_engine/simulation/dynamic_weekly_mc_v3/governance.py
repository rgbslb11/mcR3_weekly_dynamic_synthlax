from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .errors import InputValidationError


def _nonempty_rows(path: Path, sheet: str) -> list[list[Any]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    out = []
    for row in ws.iter_rows(values_only=True):
        vals = list(row)
        if any(v is not None for v in vals):
            out.append(vals)
    return out


def inspect_model_parameters(path: Path) -> dict[str, object]:
    wb = load_workbook(path, read_only=True, data_only=True)
    rules = wb["04_CCG_RULES"]
    ccg_rules: dict[str, str] = {}
    for row in rules.iter_rows(min_row=4, values_only=True):
        if row[0] is None:
            continue
        ccg_rules[str(row[0])] = str(row[2])

    params = wb["02_PARAMETER_REGISTER"]
    parameter_rows = {}
    for row in params.iter_rows(min_row=3, values_only=True):
        if row[0] is None:
            continue
        parameter_rows[str(row[0])] = row

    open_ws = wb["12_OPEN_ITEMS"]
    open_items = {}
    for row in open_ws.iter_rows(min_row=3, values_only=True):
        if row[0] is None:
            continue
        open_items[str(row[0])] = str(row[4])

    required_rules = ["R-CCG-01", "R-CCG-02", "R-CCG-05", "R-CCG-06", "R-CCG-08", "CG-8", "TB-ECL/A8"]
    missing = [r for r in required_rules if r not in ccg_rules]
    if missing:
        raise InputValidationError(f"Model Parameters missing governed CCG rules: {missing}")

    return {
        "ccg_rule_ids_present": required_rules,
        "g5_bid_rule": ccg_rules["CG-8"],
        "a8_ecl_tiebreak_rule": ccg_rules["TB-ECL/A8"],
        "aac_division_rule": ccg_rules["R-CCG-05"],
        "hfa_v2_legacy_points": float(parameter_rows["ENG-HOME-FIELD"][4]),
        "hfa_team_master_locked_points": float(parameter_rows["SCHED-HFA-BASE"][4]),
        "margin_sd_calibration_status": open_items.get("ENG-CAL-MARGIN"),
        "schedule_13_game_exception_status": open_items.get("OI-SCHED-13"),
        "aac_divisions_source_status": "RATIFIED_BUT_ROWS_NOT_MOUNTED",
    }


def inspect_bracket(path: Path) -> dict[str, object]:
    rows = _nonempty_rows(path, "Bracket Regime LOCKED")
    rules = {str(r[0]): str(r[1]) for r in rows[1:] if len(r) >= 2 and r[0] is not None}
    required = ["REGIME", "status", "AB1", "AB2", "AB3", "AB4", "S1", "S2", "S3", "S4"]
    missing = [x for x in required if x not in rules]
    if missing:
        raise InputValidationError(f"Bracket regime missing rules: {missing}")
    return {
        "status": rules["status"],
        "field_rule": rules["REGIME"],
        "ab2_original_text": rules["AB2"],
        "cg8_supersession_required": True,
        "straight_seeding": rules["S1"],
        "top4_byes": rules["S2"],
        "play_in": rules["S3"],
        "round1": rules["S4"],
        "quarterfinal_opponent_mapping_explicit": False,
    }


def inspect_playoff_calendar(path: Path) -> dict[str, object]:
    rows = _nonempty_rows(path, "Playoff_Calendar")
    by_session = {str(r[0]): r for r in rows[1:] if r[0] is not None}
    required = ["CONF CHAMPIONSHIPS", "SELECTION DAY", "ARMY–NAVY", "PLAY-IN G1", "PLAY-IN G2", "NATIONAL CHAMPIONSHIP"]
    missing = [x for x in required if x not in by_session]
    if missing:
        raise InputValidationError(f"Playoff calendar missing sessions: {missing}")
    return {
        "selection_day": str(by_session["SELECTION DAY"][3]),
        "army_navy_date": str(by_session["ARMY–NAVY"][3]),
        "army_navy_counts_to_record": "Counts to season record" in str(by_session["ARMY–NAVY"][7]),
        "army_navy_after_selection": True,
        "quarterfinal_mapping_is_generic": True,
    }


def inspect_fcs_authority(path: Path) -> dict[str, object]:
    rows = _nonempty_rows(path, "Build Manifest")
    fields = {str(r[0]): r[1] for r in rows if len(r) >= 2 and r[0] is not None}
    return {
        "ruling_applied": fields.get("Ruling applied"),
        "model_use_authorized": str(fields.get("Model use authorized", "")).upper() == "TRUE",
        "model_use_authorized_raw": fields.get("Model use authorized"),
    }


def governance_blockers(
    *,
    model_parameters: dict[str, object],
    bracket: dict[str, object],
    playoff_calendar: dict[str, object],
    fcs: dict[str, object],
    hfa_ruling_applied: bool = False,
    fcs_ruling_applied: bool = False,
    thirteen_game_exceptions_validated: bool = False,
    a8_ecl_ordering_resolved: bool = False,
    sos_semantics_governed: bool = False,
    fcs_unified_scale_governed: bool = False,
) -> list[str]:
    """Governed-evidence blockers.

    Each ``*_applied`` / ``*_validated`` flag is supplied by the caller only
    after the corresponding R2 ruling has been encoded **and** its deterministic
    validation has passed. A ruling on its own never clears anything here; the
    default for every flag is the pre-ruling, fail-closed state.
    """
    blockers: list[str] = []

    # R2-HFA-3P5 settles which value V3 uses. The registers still disagree and
    # are not edited; the ruling is what makes the disagreement non-blocking.
    if model_parameters["hfa_v2_legacy_points"] != model_parameters["hfa_team_master_locked_points"]:
        if not hfa_ruling_applied:
            blockers.append("governance.V3_HFA_BASELINE_CONFLICT_4P0_VS_3P5")

    # ENG-CAL-MARGIN stays OPEN: no ruling can substitute for calibration evidence.
    if model_parameters["margin_sd_calibration_status"] == "OPEN":
        blockers.append("governance.GAME_SD_CALIBRATION_OPEN")

    # OI-SCHED-13 stays OPEN in the register; R2-SCHED-13GAME approves the five
    # schedules and the mounted rows must validate before it clears.
    if model_parameters["schedule_13_game_exception_status"] == "OPEN":
        if not thirteen_game_exceptions_validated:
            blockers.append("governance.FIVE_13_GAME_SCHEDULE_EXCEPTIONS_UNRATIFIED")

    # The POWER_CRUNCH Build Manifest still records FALSE. R2-FCS-ELO-1250
    # supersedes it for V3 use; the manifest itself is left unedited.
    if not fcs["model_use_authorized"] and not fcs_ruling_applied:
        blockers.append("governance.FCS_SOURCE_MODEL_USE_AUTHORIZED_FALSE")

    # R2-NO-RESEED rules out reseeding and binds every edge the official artifact
    # states, but Bracket_Flow never says which R1 winner fills E/F/G/H.
    if not bracket["quarterfinal_opponent_mapping_explicit"]:
        blockers.append("governance.POSTSEASON_QUARTERFINAL_MAPPING_NOT_EXPLICIT")

    # R2-A8-ECL-ORDER breaks the cycle causally rather than by picking a tiebreak.
    if not a8_ecl_ordering_resolved:
        blockers.append("governance.A8_ECL_FINAL_BOARD_TIEBREAK_ORDERING_NOT_EXPLICIT")

    # R2-SOS fixes the weights; nothing in the repository defines the OWP/OOWP
    # denominator, exclusion or instance-weighting rules they are applied to.
    if not sos_semantics_governed:
        blockers.append("governance.OWP_OOWP_DENOMINATOR_SEMANTICS_NOT_GOVERNED")

    # R2-FCS-ELO-1250 fixes an Elo; no register maps it onto the unified
    # neutral-points axis the engine rates on.
    if not fcs_unified_scale_governed:
        blockers.append("governance.FCS_FIXED_ELO_1250_TO_UNIFIED_POINTS_SCALE_NOT_GOVERNED")

    return blockers
