from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import load_workbook


def load_v2_team_probabilities(path: Path) -> dict[str, dict[str, float | None]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Team Probabilities"]
    rows = ws.iter_rows(values_only=True)
    # Title rows precede the header; find the row beginning MC Rank.
    header = None
    for row in rows:
        if row and row[0] == "MC Rank":
            header = list(row)
            break
    if header is None:
        raise ValueError("V2.1 Team Probabilities header not found")
    idx = {str(x): i for i, x in enumerate(header) if x is not None}
    out = {}
    for row in rows:
        if not row or row[idx["Short"]] is None:
            continue
        sid = str(row[idx["Short"]])
        out[sid] = {
            "avg_wins_pre_ccg": float(row[idx["Avg Wins Pre-CCG"]]),
            "conference_title_pct": None if row[idx["Conference Title %"]] is None else float(row[idx["Conference Title %"]]),
            "cfp_pct": float(row[idx["CFP %"]]),
            "cfp_bye_pct": float(row[idx["Top-4 Bye %"]]),
            "national_title_pct": float(row[idx["National Title %"]]),
        }
    return out


def compare_summary_csv(v2_path: Path, v3_summary_csv: Path) -> list[dict[str, object]]:
    v2 = load_v2_team_probabilities(v2_path)
    with v3_summary_csv.open(newline="", encoding="utf-8") as f:
        v3 = {row["schedule_id"]: row for row in csv.DictReader(f)}
    rows = []
    for sid in sorted(v2):
        if sid not in v3:
            rows.append({"schedule_id": sid, "status": "BLOCKED_MISSING_V3_TEAM"})
            continue
        r = v3[sid]
        rows.append({
            "schedule_id": sid,
            "status": "COMPARED",
            "delta_avg_wins": float(r["average_regular_season_wins"]) - v2[sid]["avg_wins_pre_ccg"],
            "delta_conference_title_pct": None if v2[sid]["conference_title_pct"] is None else float(r["conference_title_pct"]) - float(v2[sid]["conference_title_pct"]),
            "delta_cfp_pct": float(r["cfp_pct"]) - v2[sid]["cfp_pct"],
            "delta_cfp_bye_pct": float(r["cfp_bye_pct"]) - v2[sid]["cfp_bye_pct"],
            "delta_national_title_pct": float(r["national_title_win_pct"]) - v2[sid]["national_title_pct"],
        })
    return rows
