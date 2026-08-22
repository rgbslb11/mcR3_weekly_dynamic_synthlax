"""Extract the located candidate statements of historical FBS population.

The 2024 population question was reported to this lane as a disagreement between
two numbers. Answering it means putting every located claim about 2024 in one
place with its digest attached, next to a membership derived from a season
source, and letting the reader see which claims are about real football and
which are about a synthetic universe. That is what this extract is for.

Each candidate is read once here, with its SHA-256, so the build does not depend
on staged files a reviewer may not have. Nothing extracted is treated as
authority: these are *claims*, and the extract records them as claims.

Three candidates are located and one is in-repo:

``SYNTHETIC_NCAA_CONFERENCE_DISTRIBUTION_2021_2026``
    A per-season subdivision and conference table over 251 entities. Its own
    Methodology sheet names an upstream, ``Synthetic_CFB_2021_2026_Walkover_
    RERUN.csv``, and the word in its filename is *Synthetic*.

``PHASE5J_TEAM_REGISTRY_2024``
    A registry inside a zipped research package, declaring an FBS set for 2024
    under a registry-only universe policy.

``V3_CANONICAL_TEAM_MASTER_2026``
    In-repo, and authoritative for a synthetic 2026 season. Included because its
    FBS row count is a number someone may reach for when asked how many FBS
    teams there are, and it is not an answer to that question for any real
    season.

Usage::

    python scripts/extract_population_candidates.py
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import sys
import zipfile
from pathlib import Path

import openpyxl

REPO_ROOT = Path(__file__).resolve().parents[1]
BRIDGE_ROOT = REPO_ROOT / "reference" / "dynamic_weekly_mc_v3" / "historical_bridge_r1"
EXTRACT_PATH = BRIDGE_ROOT / "mounted" / "V3_BRIDGE_POPULATION_CANDIDATE_EXTRACT.json"

DOWNLOADS = Path("C:/Users/dbaxt/Downloads")
DISTRIBUTION_PATH = DOWNLOADS / "Synthetic_NCAA_Conference_Distribution_2021_2026.xlsx"
PHASE5J_ZIP = DOWNLOADS / "Power_Crunch_Research_Lab_Phase5J_2024_FACT_Ratings_Package.zip"
CANONICAL_MASTER = (
    REPO_ROOT
    / "reference"
    / "dynamic_weekly_mc_v3"
    / "inputs"
    / "2026_TEAM_CANONICAL_MASTER_v2_LLM_GROUNDING.md"
)

_CANONICAL_ROW = re.compile(
    r"^\|\s*(\d+)\s*\|\s*`([^`]+)`\s*\|\s*([^|]+?)\s*\|\s*`([^`]+)`\s*\|"
    r"\s*`([^`]+)`\s*\|\s*`([^`]+)`\s*\|"
)


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest() if path.exists() else ""


def extract_distribution() -> dict:
    if not DISTRIBUTION_PATH.exists():
        return {"source_located": False}
    workbook = openpyxl.load_workbook(DISTRIBUTION_PATH, read_only=True, data_only=True)
    rows = list(workbook["Team Affiliations"].iter_rows(values_only=True))
    header = [str(cell) for cell in rows[0]]
    by_season: dict[str, dict[str, object]] = {}
    for season in (2021, 2022, 2023, 2024, 2025, 2026):
        column = header.index(f"{season} Subdivision")
        members = sorted(
            str(row[0]) for row in rows[1:] if row[0] and row[column] == "FBS"
        )
        counts: dict[str, int] = {}
        for row in rows[1:]:
            if not row[0]:
                continue
            label = str(row[column])
            counts[label] = counts.get(label, 0) + 1
        by_season[str(season)] = {
            "fbs_count": len(members),
            "fbs_members": members,
            "subdivision_counts": dict(sorted(counts.items())),
        }
    methodology = {
        str(row[0]): "" if len(row) < 2 or row[1] is None else str(row[1])
        for row in workbook["Methodology"].iter_rows(values_only=True)
        if row and row[0] is not None
    }
    return {
        "source_located": True,
        "candidate_id": "SYNTHETIC_NCAA_CONFERENCE_DISTRIBUTION_2021_2026",
        "source_path": str(DISTRIBUTION_PATH).replace("\\", "/"),
        "source_sha256": digest(DISTRIBUTION_PATH),
        "universe_class": "SYNTHETIC",
        "declared_upstream": methodology.get("Source", ""),
        "declared_population": methodology.get("Population", ""),
        "entities": len(rows) - 1,
        "by_season": by_season,
    }


def extract_phase5j() -> dict:
    if not PHASE5J_ZIP.exists():
        return {"source_located": False}
    with zipfile.ZipFile(PHASE5J_ZIP) as archive:
        names = [name for name in archive.namelist() if name.endswith("team_registry_2024.csv")]
        if not names:
            return {
                "source_located": False,
                "detail": "package located, team_registry_2024.csv not inside it",
            }
        member_name = names[0]
        raw = archive.read(member_name)
        rows = list(csv.DictReader(io.StringIO(raw.decode("utf-8-sig"))))
    subdivision_column = next(
        (
            column
            for column in (rows[0].keys() if rows else [])
            if column.strip().lower() in ("subdivision", "division", "level")
        ),
        "",
    )
    fbs_rows = [
        row
        for row in rows
        if not subdivision_column or str(row.get(subdivision_column, "")).upper() == "FBS"
    ]
    return {
        "source_located": True,
        "candidate_id": "PHASE5J_TEAM_REGISTRY_2024",
        "source_path": f"{PHASE5J_ZIP}!{member_name}".replace("\\", "/"),
        "package_sha256": digest(PHASE5J_ZIP),
        "member_sha256": hashlib.sha256(raw).hexdigest(),
        "universe_class": "RESEARCH_REGISTRY",
        "columns": list(rows[0].keys()) if rows else [],
        "rows": len(rows),
        "subdivision_column": subdivision_column,
        "by_season": {
            "2024": {
                "fbs_count": len(fbs_rows),
                "fbs_members": sorted(
                    str(row.get("team") or row.get("team_name") or "").strip()
                    for row in fbs_rows
                ),
            }
        },
    }


def extract_canonical_master() -> dict:
    if not CANONICAL_MASTER.exists():
        return {"source_located": False}
    entities = []
    for line in CANONICAL_MASTER.read_text(encoding="utf-8").splitlines():
        match = _CANONICAL_ROW.match(line)
        if match:
            entities.append(
                {
                    "schedule_id": match.group(2),
                    "team_name": match.group(3),
                    "entity_scope": match.group(5),
                    "division_2026": match.group(6),
                }
            )
    fbs = [row for row in entities if row["entity_scope"] == "FBS_MEMBER"]
    return {
        "source_located": True,
        "candidate_id": "V3_CANONICAL_TEAM_MASTER_2026",
        "source_path": str(CANONICAL_MASTER.relative_to(REPO_ROOT)).replace("\\", "/"),
        "source_sha256": digest(CANONICAL_MASTER),
        "universe_class": "SYNTHETIC_2026",
        "entities": len(entities),
        "by_season": {
            "2026": {
                "fbs_count": len(fbs),
                "fbs_members": sorted(row["team_name"] for row in fbs),
            }
        },
        "applies_to_historical_seasons": False,
    }


def main() -> int:
    candidates = {
        "SYNTHETIC_NCAA_CONFERENCE_DISTRIBUTION_2021_2026": extract_distribution(),
        "PHASE5J_TEAM_REGISTRY_2024": extract_phase5j(),
        "V3_CANONICAL_TEAM_MASTER_2026": extract_canonical_master(),
    }
    extract = {
        "artifact": "V3_BRIDGE_POPULATION_CANDIDATE_EXTRACT",
        "artifact_version": "V3-HISTORICAL-BRIDGE-CORPUS-R1",
        "artifact_class": "CANDIDATE_CLAIMS_NOT_AUTHORITY",
        "candidates": candidates,
    }
    EXTRACT_PATH.parent.mkdir(parents=True, exist_ok=True)
    EXTRACT_PATH.write_text(
        json.dumps(extract, indent=2, sort_keys=True, ensure_ascii=False) + "\n",
        encoding="utf-8",
        newline="\n",
    )
    print(f"wrote {EXTRACT_PATH.relative_to(REPO_ROOT)}")
    for name, record in candidates.items():
        seasons = record.get("by_season", {})
        summary = ", ".join(
            f"{season}={payload['fbs_count']}" for season, payload in sorted(seasons.items())
        )
        print(f"  {name}: located={record.get('source_located')} {summary}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
